import json, re, subprocess, sys, tempfile, unittest
from pathlib import Path
try:
    import fontTools, pptx, PIL  # noqa: F401
    HAVE=True
except ImportError: HAVE=False
from workflow import onboard, read

def make_font(path,chars='ABC 中文來源示意圖',bold=False):
    """Synthetic TrueType font: every char is a box, so tests need no installed fonts."""
    from fontTools.fontBuilder import FontBuilder
    from fontTools.pens.ttGlyphPen import TTGlyphPen
    names=['.notdef']+['u%04X'%ord(c) for c in chars]
    fb=FontBuilder(1000,isTTF=True); fb.setupGlyphOrder(names)
    fb.setupCharacterMap({ord(c):'u%04X'%ord(c) for c in chars})
    glyphs={}
    for n in names:
        pen=TTGlyphPen(None)
        if n!='u0020':
            w=120 if bold else 60
            pen.moveTo((100,0));pen.lineTo((100,700));pen.lineTo((100+w,700));pen.lineTo((100+w,0));pen.closePath()
        glyphs[n]=pen.glyph()
    fb.setupGlyf(glyphs); fb.setupHorizontalMetrics({n:(1000 if ord(chr(int(n[1:],16)))>0x2E7F else 600,0) if n.startswith('u') else (600,0) for n in names})
    fb.setupHorizontalHeader(ascent=880,descent=-120); fb.setupNameTable({'familyName':'KevinTest','styleName':'Bold' if bold else 'Regular'})
    fb.setupOS2(); fb.setupPost(); fb.save(str(path)); return path

@unittest.skipUnless(HAVE,'fonttools/python-pptx/Pillow not installed')
class VisualTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.d=Path(self.tmp.name)
        fd=self.d/'fonts';fd.mkdir();make_font(fd/'Test-Regular.ttf');make_font(fd/'Test-Bold.ttf',bold=True)
        import visuals;self.v=visuals
        self.fonts={'regular':visuals.FontStack([fd/'Test-Regular.ttf']),'bold':visuals.FontStack([fd/'Test-Bold.ttf']),'source':'test'}
    def tearDown(self):self.tmp.cleanup()
    def svg(self,body):return f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 400 100">{body}</svg>'
    def test_text_becomes_outlines(self):
        out,rep=self.v.outline_svg(self.svg('<text x="10" y="50" font-size="20">AB 中文</text><text x="10" y="90" font-size="10">來源</text>'),self.fonts)
        self.assertNotIn('<text',out);self.assertEqual(rep['texts_outlined'],2);self.assertEqual(rep['missing_glyphs'],'')
        self.assertIn('aria-label="AB 中文"',out);self.assertIn('xlink:href="#g',out)
        self.assertEqual(rep['unique_glyphs'],6)  # A B 中 文 來 源 each defined once
        self.assertEqual(self.v.check_svg(out,illustrative=False)['status'],'passed')
    def test_anchor_and_tspan(self):
        start,_=self.v.outline_svg(self.svg('<text x="200" y="50" font-size="10">中</text>'),self.fonts)
        mid,_=self.v.outline_svg(self.svg('<text x="200" y="50" font-size="10" text-anchor="middle">中</text>'),self.fonts)
        x=lambda s:float(re.search(r'translate\(([\d.]+)',s)[1])
        self.assertAlmostEqual(x(start)-x(mid),5.0,places=1)
        t,rep=self.v.outline_svg(self.svg('<text x="0" y="50" font-size="10" fill="#111111">A<tspan fill="#ff0000" font-weight="700">B</tspan>C</text>'),self.fonts)
        self.assertEqual(rep['texts_outlined'],1);self.assertIn('fill="#ff0000"',t);self.assertEqual(t.count('<use'),3)
    def test_missing_glyph_reported_and_check_flags_live_text(self):
        _,rep=self.v.outline_svg(self.svg('<text x="0" y="50">Z</text>'),self.fonts)
        self.assertEqual(rep['missing_glyphs'],'Z')
        r=self.v.check_svg(self.svg('<text x="0" y="50">來源</text>'))
        self.assertEqual(r['status'],'failed');self.assertTrue(any('<text>' in i for i in r['issues']))
        self.assertIn('缺資料來源',' '.join(self.v.check_svg(self.svg('<rect/>'))['issues']))
    def test_outline_file_keeps_live_original_and_skips_second_run(self):
        src=self.d/'01_x.svg';src.write_text(self.svg('<text x="0" y="50">來源</text>'),encoding='utf-8')
        r=self.v.outline_file(src,None,self.fonts,None,self.d/'原始碼')
        self.assertNotIn('<text',src.read_text(encoding='utf-8'));self.assertIn('<text',(self.d/'原始碼/01_x_活字原檔.svg').read_text(encoding='utf-8'))
        self.assertEqual(self.v.outline_file(src,None,self.fonts,None,self.d/'原始碼')['skipped'],'already outlined')
    def spec(self,**kw):
        return {'id':'01','name':'測試','title':'AB','source':'ABC 年報 p.1','elements':[
            {'type':'box','id':'a','x':10,'y':150,'w':300,'h':100,'title':'A','lines':['B']},
            {'type':'box','id':'b','x':500,'y':150,'w':300,'h':100,'title':'C'},{'type':'arrow','from':'a','to':'b','label':'A'}],**kw}
    def test_spec_rules(self):
        with self.assertRaises(ValueError):self.v.build_svg(self.spec(source='ABC'))
        bad=self.spec();bad['elements'].append({'type':'arrow','from':'a','to':'zz'})
        with self.assertRaises(ValueError):self.v.build_svg(bad)
        with self.assertRaises(ValueError):self.v.build_svg(self.spec(elements=[{'type':'star'}]))
        live=self.v.build_svg(self.spec(),{'name':'中文','ticker':'1234'})
        self.assertIn('示意圖',live);self.assertIn('來源：ABC 年報 p.1',live);self.assertIn('x1="310.0"',live)
    def test_produce_writes_outlined_svg_spec_and_png(self):
        spec=self.spec(title='AB 中文')
        with self.assertRaises(ValueError):self.v.produce(spec,self.d/'out',None,self.fonts)  # footer glyphs missing → refuse, no tofu
        chars=''.join(sorted(set(''.join(re.findall(r'>([^<]+)<',self.v.build_svg(spec))))))
        fd=self.d/'full';fd.mkdir();make_font(fd/'F-Regular.ttf',chars);make_font(fd/'F-Bold.ttf',chars,True)
        r=self.v.produce(spec,self.d/'out',None,self.v.load_fonts(str(fd)))
        self.assertEqual(r['qa']['status'],'passed');self.assertNotIn('<text',Path(r['svg']).read_text(encoding='utf-8'))
        self.assertEqual(json.loads(Path(r['spec']).read_text(encoding='utf-8'))['id'],'01')
        if r.get('png_error') is None:self.assertTrue(Path(r['png']).exists())

@unittest.skipUnless(HAVE,'python-pptx/Pillow not installed')
class DeckTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)/'TWSE-2330'
        self.c=onboard(self.root,'TWSE','2330','Fixture','semiconductor')
        from PIL import Image
        Image.new('RGB',(1600,900),'white').save(self.root/'工程圖解/PNG預覽/01_x.png')
        import deck;self.deck=deck
    def tearDown(self):self.tmp.cleanup()
    def spec(self):
        return {'topic':'測試','asof':'2026-10-02','default_source':'Fixture 年報 p.1','style':{'palette':'07','font':'C'},'slides':[
            {'type':'title','title':'Fixture 研究','subtitle':'副標'},{'type':'section','title':'公司'},
            {'type':'bullets','title':'結論','bullets':['一',['子項']],'takeaway':'重點'},
            {'type':'cards','title':'產品','cards':[{'title':'A','body':['x','y']},{'title':'B','body':'z'}]},
            {'type':'kpi','title':'數字','stats':[{'value':'NT$1.0億','label':'營收'}]},
            {'type':'image','title':'圖','image':'工程圖解/PNG預覽/01_x.png','caption':'示意圖'},
            {'type':'chart','title':'營收','chart':'column','categories':['2024','2025'],'series':[{'name':'營收','values':[900,1000]}]},
            {'type':'chart','title':'客戶','chart':'doughnut','categories':['A','B'],'series':[{'name':'占比','values':[0.6,0.4]}],'number_format':'0%'},
            {'type':'table','title':'同業','columns':['公司','P/E'],'rows':[['X','N/A']]},
            {'type':'closing','title':'追蹤','bullets':['催化']}]}
    def test_build_and_check(self):
        r=self.deck.build_deck(self.spec(),self.root/'整合簡報',self.c,self.root)
        self.assertEqual(r['slides'],10);self.assertEqual(r['check']['status'],'passed',r['check'])
        self.assertTrue(r['file'].endswith('Fixture2330_測試_10頁_20261002.pptx'))
        self.assertIn('Microsoft JhengHei',r['check']['fonts']);self.assertIn('Impact',r['check']['fonts'])
        from pptx import Presentation
        prs=Presentation(r['file'])
        charts=[sh.chart for s in prs.slides for sh in s.shapes if sh.has_chart]
        self.assertEqual(charts[0].value_axis.minimum_scale,0)
        theme=prs.slide_master.part.part_related_by('http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme').blob.decode()
        self.assertIn('<a:latin typeface="Impact"',theme);self.assertIn('<a:ea typeface="Microsoft JhengHei"',theme)
    def test_rules(self):
        s=self.spec();s.pop('style')
        with self.assertRaises(ValueError):self.deck.build_deck(s,self.root/'整合簡報',self.c,self.root)  # no deck_style
        s=self.spec();s['slides']=s['slides'][:-1]
        with self.assertRaises(ValueError):self.deck.build_deck(s,self.root/'整合簡報',self.c,self.root)  # closing required
        s=self.spec();s['slides'][6]['series'][0]['values']=[1]
        with self.assertRaises(ValueError):self.deck.build_deck(s,self.root/'整合簡報',self.c,self.root)
        s=self.spec();s.pop('default_source')
        with self.assertRaises(ValueError):self.deck.build_deck(s,self.root/'整合簡報',self.c,self.root)
        s=self.spec();s['slides'][8]['rows']=[['X','⚫ 478']]
        r=self.deck.build_deck(s,self.root/'整合簡報',self.c,self.root);self.assertTrue(any('符號' in w for w in r['check']['warnings']))
        s=self.spec();s['slides'][-1]['disclaimer']='研究整理'
        r=self.deck.build_deck(s,self.root/'整合簡報',self.c,self.root);self.assertEqual(r['check']['status'],'failed')
    def test_set_style_cli_and_fonts(self):
        wf=str(Path(__file__).with_name('workflow.py'))
        bad=subprocess.run([sys.executable,wf,'--root',str(self.root),'set-deck-style','--palette','11','--font','C'],capture_output=True,text=True)
        self.assertNotEqual(bad.returncode,0)
        subprocess.run([sys.executable,wf,'--root',str(self.root),'set-deck-style','--palette','05','--font','B','--cjk-font','Noto Sans TC'],check=True,capture_output=True)
        self.assertEqual(read(self.root/'config.json')['deck_style'],{'palette':'05','font':'B','cjk_font':'Noto Sans TC'})
        s=self.spec();s.pop('style');r=self.deck.build_deck(s,self.root/'整合簡報',read(self.root/'config.json'),self.root)
        self.assertEqual(r['style']['cjk'],'Noto Sans TC')
        fixed=self.deck.apply_fonts(r['file'],self.deck.resolve_style({'style':{'palette':'05','font':'B'}}))
        self.assertGreater(fixed['runs_updated'],0);self.assertIn('Microsoft JhengHei',self.deck.check_deck(r['file'])['fonts'])

if __name__=='__main__':unittest.main()
