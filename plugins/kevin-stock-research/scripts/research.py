"""Kevin research primitives. Immutable evidence, explicit gaps, no trading connection."""
from __future__ import annotations
import argparse, contextlib, csv, hashlib, io, json, math, os, re, sqlite3, subprocess, tempfile
import urllib.request, urllib.parse
from datetime import datetime, date, timezone, timedelta
from pathlib import Path

VERSION = '0.3.0'
TZ = timezone(timedelta(hours=8))
def now(): return datetime.now(TZ).isoformat(timespec='seconds')
def digest(b): return hashlib.sha256(b).hexdigest()
def dump(p, obj):
 p=Path(p); p.parent.mkdir(parents=True,exist_ok=True)
 with tempfile.NamedTemporaryFile(mode='w',encoding='utf-8',dir=p.parent,prefix=p.name+'.',suffix='.tmp',delete=False) as f:
  json.dump(obj,f,ensure_ascii=False,indent=2,default=str); t=Path(f.name)
 try: os.replace(t,p)
 finally:
  if t.exists(): t.unlink()
def read(p): return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def number(v):
 if v is None or isinstance(v,bool): return None
 try:
  x=float(str(v).replace(',','').strip()); return x if math.isfinite(x) else None
 except (ValueError,TypeError): return None
def identity(market,ticker):
 pattern=r'\d{4,6}' if market in ('TWSE','TPEX') else r'[A-Z][A-Z0-9.-]{0,9}'
 if market not in ('TWSE','TPEX','NYSE','NASDAQ') or not re.fullmatch(pattern,str(ticker)): raise ValueError('invalid market/ticker')
 return market+':'+str(ticker)
def period(year,quarter=None,month=None,basis='annual'):
 year=int(year); year=year+1911 if year<1911 else year
 if not 1900<=year<=2200: raise ValueError('invalid fiscal year')
 if quarter is not None and int(quarter) not in range(1,5): raise ValueError('invalid quarter')
 if month is not None and int(month) not in range(1,13): raise ValueError('invalid month')
 if basis not in ('annual','ytd','quarter','month'): raise ValueError('invalid period basis')
 return {'year':year,'quarter':quarter,'month':month,'basis':basis}

# Research root layout. 分類PDF holds readable hard links over 原始文件; see library.py.
LAYOUT=('原始文件','文件文字','財務數據','模型','盤後紀錄','正式成果',
 '分類PDF/01_年報','分類PDF/02_年度財報','分類PDF/03_季報','分類PDF/04_法說會資料',
 '分類PDF/05_技術簡報','分類PDF/06_券商研究','分類PDF/07_產業資料','分類PDF/08_股東會與公司治理',
 '工程圖解/PNG預覽','工程圖解/實物設備詳解版','工程圖解/工藝精度圖例增補版','整合簡報')

class Store:
 def __init__(self,root):
  self.root=Path(root); self.root.mkdir(parents=True,exist_ok=True)
  for f in LAYOUT: (self.root/f).mkdir(parents=True,exist_ok=True)
  self.db=sqlite3.connect(self.root/'research.sqlite',timeout=30)
  self.db.row_factory=sqlite3.Row
  self.db.executescript('''CREATE TABLE IF NOT EXISTS evidence(
   id TEXT PRIMARY KEY, identity TEXT, category TEXT, url TEXT, title TEXT,
   published TEXT, period TEXT, acquired TEXT, status TEXT, sha256 TEXT, path TEXT, note TEXT);
   CREATE TABLE IF NOT EXISTS snapshots(id TEXT PRIMARY KEY, identity TEXT, kind TEXT, data_date TEXT,
   acquired TEXT, sha256 TEXT, payload TEXT);
   CREATE TABLE IF NOT EXISTS alerts(id TEXT PRIMARY KEY, identity TEXT, kind TEXT, payload TEXT,
   created TEXT, delivered TEXT);
   CREATE TABLE IF NOT EXISTS fills(id TEXT PRIMARY KEY, identity TEXT, payload TEXT, created TEXT);''')
 def check_security(self,security):
  parts=security.split(':')
  if len(parts)!=2 or identity(*parts)!=security: raise ValueError('invalid security')
  config=self.root/'config.json'
  if config.exists() and read(config).get('security')!=security: raise ValueError('research root belongs to another security')
 def ingest(self,blob,*,security,category,url,title='',published=None,fiscal_period=None,status='full',suffix='.bin',note=''):
  self.check_security(security)
  if status not in ('full','partial','metadata','blocked','missing'): raise ValueError('invalid access status')
  if status=='full' and (not blob or not blob.strip()): raise ValueError('full requires nonempty content')
  if blob and suffix.lower()=='.pdf' and not blob.lstrip().startswith(b'%PDF-'): raise ValueError('response is not PDF')
  if blob and suffix.lower()=='.pdf':
   import pypdfium2 as pdfium
   doc=pdfium.PdfDocument(blob)
   try:
    if len(doc)<1: raise ValueError('PDF has no pages')
   finally: doc.close()
  if blob and suffix.lower()=='.json': json.loads(blob)
  sha=digest(blob) if blob else ''
  safe_suffix=suffix if re.fullmatch(r'\.[A-Za-z0-9]{1,8}',suffix) else '.bin'
  rel=Path('原始文件')/(sha+safe_suffix) if blob else None
  if rel and not (self.root/rel).exists():
   with (self.root/rel).open('xb') as f: f.write(blob)
  key=digest(json.dumps([security,category,url,sha,status,published,fiscal_period],sort_keys=True,ensure_ascii=False).encode())
  row=(key,security,category,url,title,published,json.dumps(fiscal_period,ensure_ascii=False) if fiscal_period else None,now(),status,sha,str(rel) if rel else '',note)
  with self.db: self.db.execute('INSERT OR IGNORE INTO evidence VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',row)
  return key
 def fetch(self,url,*,security,category,title='',published=None,fiscal_period=None,suffix=None):
  if urllib.parse.urlsplit(url).scheme!='https': raise ValueError('HTTPS required')
  try:
   req=urllib.request.Request(url,headers={'User-Agent':'KevinResearch/'+VERSION+' (+personal research)'})
   with urllib.request.urlopen(req,timeout=40) as r:
    b=r.read(80*1024*1024+1)
    if len(b)>80*1024*1024: raise ValueError('document exceeds 80MB limit')
    if not b.strip(): raise ValueError('empty response')
    if '查詢過量' in b.decode('cp950',errors='ignore'): raise ValueError('source rate limit; defer retry')
    ext=suffix or ('.pdf' if b.lstrip().startswith(b'%PDF-') else '.html')
   access='partial' if ext.lower()=='.html' else 'full'
   return self.ingest(b,security=security,category=category,url=url,title=title,published=published,fiscal_period=fiscal_period,suffix=ext,status=access,note='HTML requires content review before full status' if access=='partial' else '')
  except Exception as e:
   return self.ingest(None,security=security,category=category,url=url,title=title,status='blocked',note=type(e).__name__+': '+str(e)[:250])
 def evidence(self): return [dict(r) for r in self.db.execute('SELECT * FROM evidence ORDER BY category,published,title')]
 def extract_pdf(self,eid):
  import pypdfium2 as pdfium
  row=dict(self.db.execute('SELECT * FROM evidence WHERE id=?',(eid,)).fetchone())
  if not row['path'].endswith('.pdf'): return None
  p=self.root/'文件文字'/(row['sha256']+'.json')
  if p.exists(): return str(p)
  doc=pdfium.PdfDocument(str(self.root/row['path']))
  pages=[]
  for i in range(len(doc)):
   page=doc[i]; tp=page.get_textpage(); pages.append({'pdf_page':i+1,'text':tp.get_text_range()}); tp.close(); page.close()
  doc.close()
  dump(p,{'evidence_id':eid,'sha256':row['sha256'],'pages':pages,'ocr_needed':sum(len(x['text'].strip()) for x in pages)<100})
  return str(p)
 def snapshot(self,security,kind,data_date,payload):
  self.check_security(security)
  date.fromisoformat(data_date)
  body=json.dumps(payload,ensure_ascii=False,sort_keys=True); sha=digest(body.encode()); sid=digest((security+kind+data_date+sha).encode())
  with self.db: self.db.execute('INSERT OR IGNORE INTO snapshots VALUES(?,?,?,?,?,?,?)',(sid,security,kind,data_date,now(),sha,body))
  return sid
 def alert(self,security,kind,event_key,payload):
  self.check_security(security)
  # Stable semantic event key: content versions belong in event_key when a correction matters.
  aid=digest((security+'|'+kind+'|'+event_key).encode())
  with self.db: self.db.execute('INSERT OR IGNORE INTO alerts VALUES(?,?,?,?,?,NULL)',(aid,security,kind,json.dumps(payload,ensure_ascii=False),now()))
  return aid
 def pending(self,security=None):
  if security: self.check_security(security)
  query='SELECT * FROM alerts WHERE delivered IS NULL'
  return [dict(r) for r in self.db.execute(query+(' AND identity=?' if security else ''),(security,) if security else ())]
 def acknowledge(self,aid):
  with self.db: self.db.execute('UPDATE alerts SET delivered=? WHERE id=? AND delivered IS NULL',(now(),aid))
 def fill(self,record):
  self.check_security(record.get('security',''))
  required=('id','security','date','side','shares','price','fees','tax','source')
  if any(k not in record for k in required): raise ValueError('incomplete execution record')
  if record.get('confirmed') is not True or record['side'] not in ('buy','sell'): raise ValueError('actual confirmed fills only')
  date.fromisoformat(record['date'])
  if number(record['shares']) is None or int(record['shares'])!=record['shares'] or record['shares']<=0 or number(record['price']) is None or record['price']<=0: raise ValueError('invalid fill')
  if any(number(record[k]) is None or record[k]<0 for k in ('fees','tax')): raise ValueError('invalid costs')
  body=json.dumps(record,ensure_ascii=False,sort_keys=True)
  old=self.db.execute('SELECT payload FROM fills WHERE id=?',(record['id'],)).fetchone()
  if old and old[0]!=body: raise ValueError('conflicting fill ID; explicit correction required')
  with self.db: self.db.execute('INSERT OR IGNORE INTO fills VALUES(?,?,?,?)',(record['id'],record['security'],body,now()))
 def export(self):
  rows=self.evidence(); dump(self.root/'來源清單.json',rows)
  if rows:
   with (self.root/'來源清單.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,fieldnames=rows[0].keys()); w.writeheader(); w.writerows(rows)
  return rows

def single_quarter(current,previous=None):
 """Convert cumulative flow metrics ONLY; never stock variables or EPS."""
 flows={'revenue','gross_profit','operating_income','net_income','parent_net_income','operating_cash_flow','investing_cash_flow','financing_cash_flow','capex','depreciation'}
 if current.get('metric') not in flows: raise ValueError('only registered flow metrics can be subtracted')
 if current['basis']!='ytd': raise ValueError('YTD required')
 q=current['quarter']
 if q not in (1,2,3,4): raise ValueError('invalid quarter')
 if number(current.get('value')) is None: raise ValueError('missing value')
 if q==1: return current['value']
 keys=('security','year','metric','currency','unit','scope','revision_basis')
 if previous is None or previous['quarter']!=q-1 or previous['basis']!='ytd' or any(current.get(k)!=previous.get(k) for k in keys): raise ValueError('incompatible preceding cumulative period')
 if number(previous.get('value')) is None: raise ValueError('missing previous value')
 return current['value']-previous['value']

# ---- Kevin 雙層錨定模型（0.2.1）----
EQ_WARN, EQ_CHOOSE = .05, .10   # 非經常性占稅前：<5% 乾淨、5–10% 警示、≥10% 須明確選帳面或本業
ANCHOR_RULE = {'high_ratio':1.3,'low_ratio':0.8,'high_adj':0.85,'low_adj':1.15}
RENAMED = {'base_pe_low':'base_pe_p (√ growth term, P)','base_pe_high':'base_pe_q (^(2/3) growth term, Q)'}

def _num(x,k,positive=False,minimum=None):
 v=x.get(k)
 if isinstance(v,bool) or not isinstance(v,(int,float)) or not math.isfinite(v): raise ValueError(k+' must be a finite number')
 if positive and v<=0: raise ValueError(k+' must be positive')
 if minimum is not None and v<minimum: raise ValueError(k+' below minimum')
 return v

def anchor_bases(a):
 """雙層錨定：base_pe = ETF PE × 比例 × M_adj；M = 大盤PE ÷ 10年中位PE。"""
 for k in ('etf_pe','ratio_p','ratio_q','market_pe','market_median_pe'): _num(a,k,True)
 if a['ratio_p']>2 or a['ratio_q']>2: raise ValueError('anchor ratio out of range')
 rule={**ANCHOR_RULE,**(a.get('rule') or {})}
 m=a['market_pe']/a['market_median_pe']
 adj=rule['high_adj'] if m>rule['high_ratio'] else rule['low_adj'] if m<rule['low_ratio'] else 1.0
 return {'m_ratio':m,'m_adj':adj,'base_pe_p':a['etf_pe']*a['ratio_p']*adj,'base_pe_q':a['etf_pe']*a['ratio_q']*adj}

def earnings_quality(q):
 """帳面與本業淨利率（合併、未扣少數股權）。非經常性逐項列示：股票評價FVTPL、一次性處分、匯兌等。"""
 if q.get('scope')!='consolidated': raise ValueError('earnings_quality must use consolidated net income; NCI is applied via owner_ratio')
 for k in ('revenue','pretax','net_income'): _num(q,k,True)
 _num(q,'tax')
 parts=q.get('nonrecurring')
 if not isinstance(parts,dict): raise ValueError('nonrecurring must list items, e.g. {"fvtpl":0,"disposals":0,"fx":0}')
 for k in parts: _num(parts,k)
 rate=q['tax']/q['pretax']
 if not 0<=rate<1: raise ValueError('tax rate outside 0-1; provide consolidated_margin manually with an explanation')
 nr=sum(parts.values()); ratio=abs(nr)/q['pretax']
 return {'reported_margin':q['net_income']/q['revenue'],'core_margin':(q['pretax']-nr)*(1-rate)/q['revenue'],
         'nonrecurring':nr,'nonrecurring_items':dict(parts),'nonrecurring_to_pretax':ratio,'tax_rate':rate,
         'level':'clean' if ratio<EQ_WARN else 'warn' if ratio<EQ_CHOOSE else 'choose'}

def _margin(src,choice,label,warnings):
 """回傳（淨利率, 口徑, 盈餘品質）；src 只能提供 consolidated_margin 或 earnings_quality 其中之一。"""
 has_eq,has_m=src.get('earnings_quality') is not None,src.get('consolidated_margin') is not None
 if has_eq==has_m: raise ValueError(label+': provide exactly one of consolidated_margin or earnings_quality')
 if has_m:
  if src.get('margin_basis')!='consolidated': raise ValueError(label+': owner_ratio must not be applied to parent-attributable margin')
  warnings.append(label+': earnings quality not checked (non-recurring items unknown)')
  return _num(src,'consolidated_margin',True),'input',None
 eq=earnings_quality(src['earnings_quality'])
 if choice not in (None,'reported','core'): raise ValueError('margin_choice must be reported or core')
 if eq['level']=='choose' and choice is None:
  raise ValueError(f"{label}: non-recurring items are {eq['nonrecurring_to_pretax']:.1%} of pretax income (>=10%); set margin_choice to reported or core")
 choice=choice or 'reported'
 if eq['level']=='warn' and choice=='reported': warnings.append(f"{label}: non-recurring items are {eq['nonrecurring_to_pretax']:.1%} of pretax income; reported margin used")
 return eq[choice+'_margin'],choice,eq

def kevin_calculate(x):
 """Kevin 雙層錨定模型純公式。EPS 與成長率不含任何倍率；target_premium 只乘在模型價。"""
 for old,new in RENAMED.items():
  if old in x: raise ValueError(f'{old} was renamed to {new} in 0.2.1')
 if 'multiplier' in x and x['multiplier']!=1:
  raise ValueError('multiplier on EPS was removed in 0.2.1: it also inflated growth and P/E (x1.5 became x1.92 on the target). Use target_premium, which applies to model_value only')
 for k in ('revenue_ytd','consolidated_margin','owner_ratio','capital_thousands','par','previous_eps','base_pe_p','base_pe_q'): _num(x,k,True)
 months=_num(x,'months')
 if not 1<=months<=12 or int(months)!=months: raise ValueError('months out of range')
 if x.get('margin_basis')!='consolidated': raise ValueError('owner_ratio must not be applied to parent-attributable margin')
 if x['owner_ratio']>1 and not x.get('nci_loss_verified'): raise ValueError('ownership ratio requires verification of negative NCI')
 premium=x.get('target_premium',1)
 if isinstance(premium,bool) or not isinstance(premium,(int,float)) or not 0<premium<=3: raise ValueError('target_premium must be in (0,3]')
 annual=_num(x,'annual_revenue',True) if 'annual_revenue' in x else x['revenue_ytd']*12/months
 eps=annual*x['consolidated_margin']*x['owner_ratio']/x['capital_thousands']*x['par']
 g=eps/x['previous_eps']-1
 if g<0: raise ValueError('original fractional-power model not applicable to negative growth')
 pe_p=x['base_pe_p']+math.sqrt(g*100); pe_q=x['base_pe_q']+(g*100)**(2/3); value=eps*(pe_p+pe_q)/2
 out={'annual_revenue':annual,'eps':eps,'growth':g,'pe_p':pe_p,'pe_q':pe_q,'price_p':eps*pe_p,'price_q':eps*pe_q,'model_value':value,'target_premium':premium}
 if premium!=1: out['model_value_with_premium']=value*premium
 return out

def kevin_valuation(x):
 """完整 Kevin 估值：盈餘品質、基期 EPS 口徑、情境A（線性年化），選用情境B（近月動能 × 最近一季淨利率）。"""
 warnings=[]; flags=[]; choice=x.get('margin_choice')
 margin,source,eq=_margin(x,choice,'scenario A',warnings)
 basis=x.get('previous_eps_basis','reported')
 if basis not in ('reported','normalized'): raise ValueError('previous_eps_basis must be reported or normalized')
 extra={}
 if basis=='normalized':
  extra['previous_eps_reported']=_num(x,'previous_eps_reported')
  if not str(x.get('previous_eps_note') or '').strip(): raise ValueError('normalized previous_eps requires previous_eps_note explaining the adjustment')
  extra['previous_eps_note']=x['previous_eps_note']
 keep=('revenue_ytd','months','owner_ratio','capital_thousands','par','previous_eps','base_pe_p','base_pe_q','nci_loss_verified','target_premium','multiplier')+tuple(RENAMED)
 core={k:x[k] for k in keep if k in x}
 a=kevin_calculate({**core,'consolidated_margin':margin,'margin_basis':'consolidated'})
 a.update(margin=margin,margin_source=source)
 if basis=='normalized' and extra['previous_eps_reported']>0: a['growth_vs_reported_eps']=a['eps']/extra['previous_eps_reported']-1
 result={'model':'kevin_dual_anchor','primary':'A','scenarios':{'A':a},'earnings_quality':eq,'previous_eps_basis':basis,**extra,
         'model_value':a['model_value'],'warnings':warnings,'flags':flags}
 rr=x.get('run_rate')
 if rr is None: return result
 if x['months']>=12: raise ValueError('run_rate needs months below 12')
 recent=rr.get('recent_months')
 if not isinstance(recent,list) or len(recent)<2: raise ValueError('run_rate.recent_months needs at least two monthly revenues')
 for v in recent: _num({'recent_month':v},'recent_month',True)
 avg=sum(recent)/len(recent); annual=x['revenue_ytd']+avg*(12-x['months'])
 margin_b,source_b,eq_b=_margin(rr,rr.get('margin_choice',choice),'scenario B',warnings)
 try:
  b=kevin_calculate({**core,'consolidated_margin':margin_b,'margin_basis':'consolidated','annual_revenue':annual})
  b.update(margin=margin_b,margin_source=source_b,recent_month_average=avg,margin_period=rr.get('margin_period'))
 except ValueError as e: b={'status':'not_applicable','reason':str(e)}
 if rr.get('pull_in_suspected'): flags.append('scenario B: customer pull-in suspected; run-rate may not persist')
 if rr.get('inorganic_monthly') is not None:
  inorganic=_num(rr,'inorganic_monthly',minimum=0); b['inorganic_share_of_recent']=inorganic/avg
  if inorganic>0: flags.append(f'scenario B: {inorganic/avg:.1%} of recent monthly revenue comes from acquisitions')
 result['scenarios']['B']=b; result['earnings_quality_B']=eq_b
 if 'model_value' in b: result['b_vs_a']=b['model_value']/a['model_value']-1
 return result

def size_position(p,signal):
 keys=('investable_assets','available_cash','holding_shares','average_cost','single_stock_limit','sector_limit','sector_value','loss_budget_remaining','risk_per_share','core_ratio','tactical_ratio','batch_budget','fee_rate','minimum_fee','sell_fee_rate','sell_tax_rate')
 missing=[k for k in keys if number(p.get(k)) is None]
 if missing: return {'status':'conditional_only','missing':missing,'shares':None}
 if not signal.get('model_verified') or not signal.get('data_current') or not signal.get('thesis_valid'): return {'status':'blocked','shares':None,'reason':'model/data/thesis gate'}
 price=number(signal.get('price')); buy_max=number(signal.get('buy_max'))
 if price is None or price<=0 or buy_max is None: return {'status':'conditional_only','shares':None,'reason':'price/buy zone missing'}
 if any(p[k]<0 for k in keys) or any(not 0<=p[k]<=1 for k in ('single_stock_limit','sector_limit','core_ratio','tactical_ratio','fee_rate','sell_fee_rate','sell_tax_rate')) or abs(p['core_ratio']+p['tactical_ratio']-1)>1e-8 or p['risk_per_share']<=0 or p['investable_assets']<=0: raise ValueError('invalid personal risk settings')
 if p['holding_shares']!=int(p['holding_shares']) or p['sector_value']+1e-6<p['holding_shares']*price: raise ValueError('inconsistent holdings/sector exposure')
 if price>buy_max: return {'status':'observe','shares':0,'reason':'outside user-defined buy zone'}
 stock_room=max(0,p['investable_assets']*p['single_stock_limit']-p['holding_shares']*price)
 sector_room=max(0,p['investable_assets']*p['sector_limit']-p['sector_value'])
 loss_unit=p['risk_per_share']+price*(p['fee_rate']+p['sell_fee_rate']+p['sell_tax_rate'])
 limits=[int(max(0,p['available_cash']-p['minimum_fee'])/(price*(1+p['fee_rate']))),int(max(0,p['batch_budget']-p['minimum_fee'])/(price*(1+p['fee_rate']))),int(stock_room/price),int(sector_room/price),int(max(0,p['loss_budget_remaining']-2*p['minimum_fee'])/loss_unit)]
 n=max(0,min(limits)); fees=max(p['minimum_fee'],n*price*p['fee_rate']) if n else 0
 return {'status':'proposal_only','shares':n,'amount':n*price+fees,'fees':fees,'currency':'NT$','limiting_shares':dict(zip(('cash','batch','stock','sector','loss'),limits)),'stock_exposure_after':(p['holding_shares']+n)*price/p['investable_assets'],'sector_exposure_after':(p['sector_value']+n*price)/p['investable_assets'],'core_ratio':p['core_ratio'],'tactical_ratio':p['tactical_ratio'],'holdings_updated':False}

def ledger(fills,price):
 securities={f.get('security') for f in fills}
 if len(securities)>1: raise ValueError('ledger must contain one security only')
 qty=0; cost=0.; realized=0.; charges=0.; contributed=0.
 for f in sorted(fills,key=lambda x:(x['date'],x['id'])):
  n=f['shares']; gross=n*f['price']; fees=f['fees']+f['tax']; charges+=fees
  if f['side']=='buy': qty+=n; cost+=gross+fees; contributed+=gross+fees
  else:
   if n>qty: raise ValueError('sell exceeds documented holdings; opening balance required')
   removed=cost*n/qty; realized+=gross-fees-removed; qty-=n; cost-=removed
 unrealized=qty*price-cost
 return {'shares':qty,'average_cost':cost/qty if qty else None,'realized_pnl':realized,'unrealized_pnl':unrealized,'total_pnl':realized+unrealized,'total_return_on_gross_purchases':(realized+unrealized)/contributed if contributed else None,'fees_and_tax':charges,'note':'Return denominator is gross documented purchases; not time-weighted return.'}

def parse_twse(j,ticker,expected_date):
 if j.get('stat')!='OK': raise ValueError('source unavailable or no data; not proof of holiday')
 if j.get('date') and str(j['date'])!=expected_date.replace('-',''): raise ValueError('stale or mismatched source date')
 tables=j.get('tables') or [j]
 for table in tables:
  fields=table.get('fields',[])
  if len(set(fields))!=len(fields):
   groups=[g['title'] for g in table.get('groups',[]) for _ in range(g['span'])]
   if len(groups)!=len(fields): raise ValueError('ambiguous duplicate field names')
   fields=[(groups[i]+'_'+f) if groups[i] else f for i,f in enumerate(fields)]
   if len(set(fields))!=len(fields): raise ValueError('non-unique grouped fields')
  for row in table.get('data',[]):
   if len(row)==len(fields) and str(row[0]).strip()==str(ticker): return dict(zip(fields,row))
 raise ValueError('ticker not present in response')

def daily(store,config,asof):
 security=identity(config['market'],config['ticker']); day=date.fromisoformat(asof)
 if asof>datetime.now(TZ).date().isoformat(): raise ValueError('future date')
 closed=config.get('confirmed_closed_dates',[])
 if not config.get('_calendar_verified_open') and (day.weekday()>=5 or asof in closed): return {'status':'closed','data_date':asof,'evidence':'weekend or configured official calendar','daily_updated':False}
 if config['market']!='TWSE': return {'status':'blocked','reason':'TPEX adapter needs source verification; do not use TWSE for TPEX'}
 d=asof.replace('-',''); t=config['ticker']; base='https://www.twse.com.tw/rwd/zh/'
 urls={
 'price':base+f'afterTrading/STOCK_DAY?date={d}&stockNo={t}&response=json',
 'institutional':base+f'fund/T86?date={d}&selectType=ALLBUT0999&response=json',
 'margin':base+f'marginTrading/MI_MARGN?date={d}&selectType=ALL&response=json',
 'short_borrow':base+f'marginTrading/TWT93U?date={d}&response=json'}
 out={'security':security,'requested_date':asof,'acquired':now(),'status':'complete','items':{},'gaps':[]}
 for kind,url in urls.items():
  eid=store.fetch(url,security=security,category='market_'+kind,suffix='.json',title=asof+' '+kind)
  row=store.db.execute('SELECT * FROM evidence WHERE id=?',(eid,)).fetchone()
  try:
   if row['status']!='full': raise ValueError(row['note'])
   j=read(store.root/row['path'])
   if kind=='price':
    if j.get('stat')!='OK': raise ValueError('price source not OK')
    target=f'{day.year-1911:03d}/{day.month:02d}/{day.day:02d}'
    matches=[r for r in j.get('data',[]) if r[0]==target]
    if len(matches)!=1: raise ValueError('exact closing date missing; calendar not inferred')
    result=dict(zip(j['fields'],matches[0]))
   else: result=parse_twse(j,t,asof)
   out['items'][kind]={'data_date':asof,'source_id':eid,'raw_fields':result}
   store.snapshot(security,kind,asof,result)
  except Exception as e: out['gaps'].append({'kind':kind,'reason':str(e),'source_id':eid})
 if out['gaps']: out['status']='incomplete'
 for kind,url,ext in [('announcements','https://openapi.twse.com.tw/v1/opendata/t187ap04_L','.json'),('tdcc_weekly','https://opendata.tdcc.com.tw/getOD.ashx?id=1-5','.csv')]:
  eid=store.fetch(url,security=security,category=kind,suffix=ext,title=kind+' latest')
  er=store.db.execute('SELECT * FROM evidence WHERE id=?',(eid,)).fetchone()
  try:
   if er['status']!='full': raise ValueError(er['note'])
   raw=(store.root/er['path']).read_bytes()
   if kind=='announcements':
    allrows=json.loads(raw)
    if not allrows or not isinstance(allrows,list): raise ValueError('announcement source empty')
    dates={roc_date(r['出表日期']) for r in allrows}
    if dates!={asof}: raise ValueError('announcement export date not current; no no-news inference')
    items=[r for r in allrows if r.get('公司代號')==str(t) and roc_date(r['發言日期'])<=asof]
    result={'source_export_date':asof,'announcements':items,'coverage':'current public snapshot; archive each run, not full historical feed'}; data_date=asof
   else:
    rows=list(csv.DictReader(io.StringIO(raw.decode('utf-8-sig')))); items=[r for r in rows if r.get('證券代號','').strip()==str(t)]
    dates={roc_date(r['資料日期']) for r in items}
    if len(dates)!=1: raise ValueError('TDCC week missing or ambiguous')
    data_date=dates.pop()
    if data_date>asof or (day-date.fromisoformat(data_date)).days>10: raise ValueError('TDCC future or excessively stale week')
    result={'data_date':data_date,'rows':items,'unit':'shares','interpretation':'holding bands; not identified major investors'}
   out['items'][kind]={'data_date':data_date,'source_id':eid,'raw_fields':result}; store.snapshot(security,kind,data_date,result)
  except Exception as e: out['gaps'].append({'kind':kind,'reason':str(e),'source_id':eid})
 out['status']='incomplete' if out['gaps'] else 'complete'
 store.snapshot(security,'daily_run',asof,out)
 folder=store.root/'盤後紀錄'/asof; folder.mkdir(parents=True,exist_ok=True)
 dump(folder/(digest(json.dumps(out,sort_keys=True).encode())[:12]+'.json'),out)
 store.export(); return out

def roc_date(value):
 s=re.sub(r'[^0-9]','',str(value))
 if len(s)==7: s=str(int(s[:3])+1911)+s[3:]
 if len(s)!=8: raise ValueError('unrecognized date')
 return date(int(s[:4]),int(s[4:6]),int(s[6:8])).isoformat()

def rolling(rows,field,n,expected_dates=None):
 if not expected_dates or n<=0: return None
 if expected_dates!=sorted(set(expected_dates)): raise ValueError('calendar must be unique and ordered')
 for d in expected_dates: date.fromisoformat(d)
 required=expected_dates[-n:]
 if len(required)!=n: return None
 unique={}
 for row in rows:
  key=row['data_date']
  if key in unique and unique[key]!=row: raise ValueError('select an explicit source revision before rolling')
  unique[key]=row
 if any(d not in unique for d in required): return None
 seq=[unique[d] for d in required]
 if any(number(x.get(field)) is None for x in seq): return None
 return sum(number(x[field]) for x in seq)

def audit_model(path,output,ticker):
 import openpyxl, zipfile
 path=Path(path); before=digest(path.read_bytes())
 with zipfile.ZipFile(path) as z:
  if any('externalLink' in n or 'vbaProject' in n for n in z.namelist()): raise ValueError('external links/macros require separate review')
 w=openpyxl.load_workbook(path); v=openpyxl.load_workbook(path,data_only=True)
 sheets=[]; errors=[]; targets=[]
 for s in w:
  for row in s:
   for c in row:
    if v[s.title][c.coordinate].data_type=='e': errors.append([s.title,c.coordinate,v[s.title][c.coordinate].value])
   if any(str(c.value)==str(ticker) for c in row): targets.append({'sheet':s.title,'row':row[0].row,'cells':{c.coordinate:{'formula_or_value':c.value,'cached':v[s.title][c.coordinate].value} for c in row}})
  sheets.append({'sheet':s.title,'rows':s.max_row,'columns':s.max_column})
 formula_count=sum(c.data_type=='f' for s in w for row in s for c in row)
 cached_count=sum(c.data_type=='f' and v[s.title][c.coordinate].value is not None for s in w for row in s for c in row)
 result={'sha256':before,'sheets':sheets,'targets':targets,'formula_count':formula_count,'cached_formula_count':cached_count,'existing_errors':errors,'status':'needs_recalculation_and_financial_validation','original_unchanged':before==digest(path.read_bytes())}
 dump(output,result); return result

def update_copy(original,destination,changes,allowlist,profile=None):
 import openpyxl,zipfile
 original=Path(original); destination=Path(destination)
 if original.resolve()==destination.resolve() or destination.exists(): raise ValueError('new destination required; never overwrite original or prior version')
 with zipfile.ZipFile(original) as archive:
  if any('externalLink' in name or 'vbaProject' in name for name in archive.namelist()): raise ValueError('external links/macros require separate review')
 before=digest(original.read_bytes()); w=openpyxl.load_workbook(original)
 from workflow import validate_mapping
 validated=validate_mapping(w,before,profile)
 if set(allowlist)!=set(validated): raise ValueError('allowlist must equal validated profile inputs')
 logs=[]
 for change in changes:
  key=change['sheet']+'!'+change['cell']
  if key not in allowlist: raise ValueError('input not approved in mapping: '+key)
  if not all(change.get(k) for k in ('source_id','asof','basis')): raise ValueError('provenance required')
  if change.get('security')!=profile['security']: raise ValueError('change security differs from mapping')
  if change.get('basis')!=validated[key]['unit']: raise ValueError('input unit differs from mapping')
  date.fromisoformat(change['asof'])
  c=w[change['sheet']][change['cell']]
  if c.data_type=='f': raise ValueError('formula edit requires separately reviewed model version')
  if isinstance(change['value'],str) and change['value'].startswith(('=','+','@')): raise ValueError('input formula injection prohibited')
  logs.append({**change,'previous':c.value}); c.value=change['value']
 destination.parent.mkdir(parents=True,exist_ok=True); w.save(destination)
 assert digest(original.read_bytes())==before
 dump(destination.with_suffix('.changes.json'),{'original_sha256':before,'changes':logs,'requires_recalculation':True})

def recalculate(source,output_dir,soffice):
 source=Path(source).resolve(); out=Path(output_dir).resolve(); out.mkdir(parents=True,exist_ok=True)
 if source.parent==out or (out/source.name).exists(): raise ValueError('recalculation requires clean separate directory')
 with tempfile.TemporaryDirectory(prefix='kevin_lo_') as profile:
  cmd=[str(soffice),'-env:UserInstallation='+Path(profile).as_uri(),'--headless','--convert-to','xlsx','--outdir',str(out),str(source)]
  r=subprocess.run(cmd,capture_output=True,text=True,errors='replace',timeout=120)
  if r.returncode or not (out/source.name).exists(): raise RuntimeError('LibreOffice recalculation failed: '+r.stderr[:200])
 return out/source.name

def main():
 # All CLI entry points use the same validated v0.2 workflow.
 from workflow import main as workflow_main
 return workflow_main()
if __name__=='__main__': main()
