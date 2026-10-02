"""分類PDF: canonical names, SHA256 dedupe, hard-link placement, index outputs and pre-download checks.

Classification is a readable view over evidence. Originals in 原始文件 are never renamed, rewritten or deleted.
"""
from __future__ import annotations
import html, os, re, shutil, urllib.parse
from datetime import date
from pathlib import Path
from research import LAYOUT, digest, dump, read, now

LIBRARY='分類PDF'
PLACEHOLDER='目前無檔案.txt'
CATEGORIES={'01_年報':'年報','02_年度財報':'年度財報','03_季報':'季報','04_法說會資料':'法說會資料',
            '05_技術簡報':'技術簡報','06_券商研究':'券商研究','07_產業資料':'產業資料','08_股東會與公司治理':'股東會與公司治理'}
DATE_MEANING={'01_年報':'會計年度','02_年度財報':'全年會計期間','03_季報':'季度及年初至當季累計',
              '04_法說會資料':'法說會日期','05_技術簡報':'文件日期','06_券商研究':'報告日期',
              '07_產業資料':'特刊年度或發布日期','08_股東會與公司治理':'股東會日期或表列資料日'}
# MOPS filename codes recognized by this release. Unknown codes stay unclassified; never guess.
MOPS_CODES={'AI1':'合併財務報告（YYYYQQ；QQ=01–03 季報，04 全年財報）','F04':'股東會年報（中文；檔名年度為會計年度）','FE4':'股東會年報（英文）'}
QUARTER_TITLE={1:'合併季報_第一季',2:'合併季報_含上半年累計',3:'合併季報_含前三季累計'}
LANGUAGES=('中文','英文')
NOTES={'02_年度財報':'全年財務報告，不是股東年報，也不是單獨Q4財報。','03_季報_累計':'Q2／Q3含累計報表；未將累計數直接標成單季數。'}
VARIANT_NOTE='同年同語言存在不同檔案指紋；以來源版區分，不推定修訂先後。'
SUFFIXES=('股份有限公司','有限公司','科技','精密','電子','工業','實業')
RULE='{公司}_{代碼}_{會計年度或文件日期}_{文件類別}_{語言}.pdf；外部研究採 {發布機構}_{日期}_{主題}_{語言}.pdf；同標題不同SHA256加 _來源版{SHA前8碼}'

def short_name(config):
    if config.get('short_name'): return config['short_name']
    name=config['name']
    for s in SUFFIXES:
        if name.endswith(s) and len(name)-len(s)>=2: return name[:-len(s)]
    return name

def original_name(row):
    """Source filename as published: MOPS keeps it in ?filename=, other sites in the URL path."""
    parts=urllib.parse.urlsplit(row.get('url') or '')
    q=urllib.parse.parse_qs(parts.query)
    if q.get('filename'): return Path(q['filename'][0]).name
    tail=Path(urllib.parse.unquote(parts.path)).name
    if tail.lower().endswith('.pdf'): return tail
    title=row.get('title') or ''
    return title if title.lower().endswith('.pdf') else ''

def parse_mops(filename):
    base=Path(filename or '').name
    m=re.match(r'^(\d{4})(0[1-4])_(\d{4,6})_AI1(?:[_.]|$)',base)
    if m:
        year,q,ticker=m[1],int(m[2]),m[3]
        if q==4: return {'ticker':ticker,'code':'AI1','category':'02_年度財報','label':year,'title':'年度合併財報','language':'中文'}
        return {'ticker':ticker,'code':'AI1','category':'03_季報','label':f'{year}Q{q}','title':QUARTER_TITLE[q],'language':'中文'}
    m=re.match(r'^(\d{4})_(\d{4,6})_(\d{8})(F04|FE4)(?:[_.]|$)',base)
    if m:
        fy,ticker,meet,code=m[1],m[2],m[3],m[4]
        out={'ticker':ticker,'code':code,'category':'01_年報','label':fy,'title':'年報','language':'英文' if code=='FE4' else '中文',
             'meeting_date':f'{meet[:4]}-{meet[4:6]}-{meet[6:]}'}
        if int(meet[:4])!=int(fy)+1: out['warning']='股東會年度不是會計年度＋1，請核對封面會計年度'
        return out
    return None

CANONICAL=re.compile(r'^(?P<publisher>[^_]+)_(?:(?P<ticker>\d{4,6})_)?(?P<label>\d{4}(?:Q[1-4]|-\d{2}-\d{2})?)_(?P<title>.+)_(?P<language>中文|英文)(?:_來源版(?P<variant>[0-9a-f]{8}))?\.pdf$')
def parse_canonical(filename):
    m=CANONICAL.match(Path(filename or '').name)
    return {k:v for k,v in m.groupdict().items() if v} if m else None

def doc_key(category,label,language,title=''):
    """Semantic identity used to avoid re-downloading the same statutory document under another filename."""
    if category=='01_年報' and title=='年報': return (category,label,language)
    if category in ('02_年度財報','03_季報') and '合併' in title: return (category,label,language)
    return None

def key_from_name(filename,ticker):
    p=parse_mops(filename)
    if p: return doc_key(p['category'],p['label'],p['language'],p['title']) if p['ticker']==str(ticker) else None
    c=parse_canonical(filename)
    if not c or c.get('ticker')!=str(ticker): return None
    title=c['title']
    category='01_年報' if title=='年報' else '02_年度財報' if title=='年度合併財報' else '03_季報' if title.startswith('合併季報') else None
    return doc_key(category,c['label'],c['language'],title) if category else None

def validate_entry(e):
    c=e.get('category')
    if c not in CATEGORIES: raise ValueError('unknown category: '+str(c))
    label=str(e.get('label') or ''); title=str(e.get('title') or '')
    if e.get('language') not in LANGUAGES: raise ValueError('language must be 中文 or 英文')
    if not title or not e.get('publisher'): raise ValueError('title and publisher required')
    if c in ('01_年報','02_年度財報') and not re.fullmatch(r'\d{4}',label): raise ValueError(c+' label must be fiscal year YYYY')
    if c=='02_年度財報' and re.search(r'Q4|第四季|第4季',title): raise ValueError('full-year report must not be named Q4')
    if c=='03_季報':
        m=re.fullmatch(r'\d{4}Q([1-4])',label)
        if not m: raise ValueError('03_季報 label must be YYYYQn')
        if m[1]=='4': raise ValueError('full-year report belongs to 02_年度財報, not Q4')
        if m[1] in '23' and '累計' not in title: raise ValueError('Q2/Q3 statements contain cumulative figures; title must say 累計')
    if c in ('04_法說會資料','06_券商研究','08_股東會與公司治理'):
        if not re.fullmatch(r'\d{4}-\d{2}-\d{2}',label): raise ValueError(c+' label must be YYYY-MM-DD')
        date.fromisoformat(label)
    if c in ('05_技術簡報','07_產業資料'):
        if not re.fullmatch(r'\d{4}(-\d{2}-\d{2})?',label): raise ValueError(c+' label must be YYYY or YYYY-MM-DD')
        if len(label)>4: date.fromisoformat(label)
    return e

def safe(part): return re.sub(r'[\\/:*?"<>|\s]+','-',str(part)).strip('-')
def build_name(e,ticker,company):
    parts=[e['publisher']]+([ticker] if e['publisher']==company else [])+[e['label'],e['title'],e['language']]
    return '_'.join(safe(p) for p in parts)+('_來源版'+e['variant'] if e.get('variant') else '')+'.pdf'

def pdf_pages(path):
    try:
        import pypdfium2 as pdfium
        doc=pdfium.PdfDocument(str(path))
        try: return len(doc)
        finally: doc.close()
    except Exception: return None

def previous_index(root):
    p=Path(root)/LIBRARY/'PDF分類索引.json'
    if not p.exists(): return []
    data=read(p)
    return data.get('files') or data.get('documents') or []

FIELDS=('category','label','title','language','publisher','note','classification_basis')
def _merge(group,item):
    for k in ('source_ids','source_urls','published_dates','original_filenames'):
        for v in item.get(k,[]):
            if v and v not in group[k]: group[k].append(v)

def collect(store,config,manifest=None):
    root=store.root; ticker=str(config['ticker'])
    groups={}
    def group(sha,path):
        g=groups.setdefault(sha,{'sha256':sha,'path':Path(path),'source_ids':[],'source_urls':[],'published_dates':[],'original_filenames':[],
                                 'placed':[],'override':{},'parsed':None,'previous':None})
        return g
    for r in store.evidence():
        if r['status']!='full' or not str(r['path']).lower().endswith('.pdf'): continue
        g=group(r['sha256'],root/r['path']); name=original_name(r)
        _merge(g,{'source_ids':[r['id']],'source_urls':[r['url']],'published_dates':[r['published']],'original_filenames':[name]})
    for p in sorted((root/LIBRARY).glob('*/*.pdf')):
        if p.parent.name not in CATEGORIES: continue
        sha=digest(p.read_bytes()); g=group(sha,p); g['placed'].append(p)
    for prev in previous_index(root):
        sha=prev.get('sha256')
        if sha in groups:
            groups[sha]['previous']=prev
            _merge(groups[sha],{'source_ids':prev.get('source_ids',[]),'source_urls':prev.get('source_urls',[]),
                                'published_dates':prev.get('published_dates',[]),'original_filenames':prev.get('original_filenames',[])})
    for item in manifest or []:
        override={k:item[k] for k in FIELDS if item.get(k) not in (None,'')}
        if item.get('local_path'):
            p=Path(item['local_path']); b=p.read_bytes()
            if not b.lstrip().startswith(b'%PDF-'): raise ValueError('manifest local_path is not PDF: '+p.name)
            g=group(digest(b),p)
            _merge(g,{'source_urls':[item.get('url')],'original_filenames':[item.get('original_filename') or p.name]})
        else:
            match=[g for g in groups.values() if (item.get('sha256') and g['sha256']==item['sha256'])
                   or (item.get('evidence_id') and item['evidence_id'] in g['source_ids'])
                   or (item.get('original_filename') and item['original_filename'] in g['original_filenames'])]
            if len(match)!=1: raise ValueError('manifest override must match exactly one PDF: '+str(item.get('sha256') or item.get('evidence_id') or item.get('original_filename')))
            g=match[0]
        g['override'].update(override)
    company=short_name(config)
    for g in groups.values():
        for name in g['original_filenames']:
            p=parse_mops(name)
            if p and p['ticker']!=ticker: g['foreign']=name
            elif p and not g['parsed']: g['parsed']=p
    return groups,company

def describe(g,company,ticker):
    base={}
    if g['placed']:
        c=parse_canonical(g['placed'][0].name)
        if c: base.update({k:c[k] for k in ('publisher','label','title','language') if k in c}); base['category']=g['placed'][0].parent.name
        base['classification_basis']='分類資料夾既有檔名'
    if g['parsed']:
        base.update({k:g['parsed'][k] for k in ('category','label','title','language')}); base['publisher']=company
        base['classification_basis']='MOPS檔名代碼 '+g['parsed']['code']
    if g['previous']:
        base.update({k:g['previous'][k] for k in FIELDS if g['previous'].get(k) not in (None,'')})
    if g['override']:
        base.update(g['override'])
        if 'classification_basis' not in g['override']: base['classification_basis']='manifest指定'
    if not base.get('category'): return None
    base.setdefault('publisher',company); base.setdefault('language','中文')
    notes=[base.get('note','')]
    if base['category']=='02_年度財報': notes.append(NOTES['02_年度財報'])
    if base['category']=='03_季報' and re.search(r'Q[23]$',str(base.get('label',''))): notes.append(NOTES['03_季報_累計'])
    if g['parsed'] and g['parsed'].get('warning'): notes.append(g['parsed']['warning'])
    base['note']='；'.join(dict.fromkeys(n for n in notes if n))
    return validate_entry(base)

def classify(store,config,manifest=None,apply=False):
    root=store.root; ticker=str(config['ticker'])
    groups,company=collect(store,config,manifest)
    files=[];unclassified=[];conflicts=[]
    for sha,g in sorted(groups.items(),key=lambda kv:(kv[1]['original_filenames'][:1],kv[0])):
        if g.get('foreign') and not g['override']:
            unclassified.append({'sha256':sha,'original_filenames':g['original_filenames'],'reason':'MOPS檔名屬於其他股票代碼：'+g['foreign']});continue
        try: e=describe(g,company,ticker)
        except ValueError as err:
            unclassified.append({'sha256':sha,'original_filenames':g['original_filenames'],'reason':str(err)});continue
        if not e:
            unclassified.append({'sha256':sha,'original_filenames':g['original_filenames'],'reason':'無法由檔名判定類別；請以manifest指定category/label/title'});continue
        placed_cats={p.parent.name for p in g['placed']}
        if placed_cats and e['category'] not in placed_cats:
            conflicts.append({'sha256':sha,'placed':[p.relative_to(root).as_posix() for p in g['placed']],'computed_category':e['category'],'action':'保留原位置，未搬移'})
        files.append({**e,'sha256':sha,'group':g})
    # Same readable name with different bytes: every new file in the group carries a short source code.
    wanted={}
    for f in files:
        if f['group']['placed']: continue
        f['base_name']=build_name({**f,'variant':None},ticker,company)
        wanted.setdefault((f['category'],f['base_name']),[]).append(f)
    occupied={(p.parent.name,p.name):sha for sha,g in groups.items() for p in g['placed']}
    for key,items in wanted.items():
        clash=len(items)>1 or key in occupied
        for f in items:
            if clash:
                f['variant']=f['sha256'][:8]; f['note']='；'.join(n for n in (f.get('note'),VARIANT_NOTE) if n)
            f['name']=build_name(f,ticker,company)
    records=[];actions=[]
    for f in files:
        g=f.pop('group')
        if g['placed']:
            target=g['placed'][0]; storage='existing'; f['category']=target.parent.name  # count where the file actually lives
        else:
            target=root/LIBRARY/f['category']/f['name']; storage='planned'
            if apply:
                target.parent.mkdir(parents=True,exist_ok=True)
                if target.exists():
                    if digest(target.read_bytes())!=f['sha256']: raise ValueError('target exists with different bytes: '+target.name)
                    storage='existing'
                else:
                    try: os.link(g['path'],target); storage='hardlink'
                    except OSError: shutil.copy2(g['path'],target); storage='copy'
            actions.append({'from':str(g['path'].relative_to(root)) if g['path'].is_relative_to(root) else g['path'].name,'to':target.relative_to(root).as_posix(),'storage':storage})
        original=root/'原始文件'/(f['sha256']+'.pdf')
        records.append({'sha256':f['sha256'],'category':f['category'],'label':f['label'],'date_meaning':DATE_MEANING[f['category']],
                        'title':f['title'],'language':f['language'],'publisher':f['publisher'],'name':target.name,
                        'pages':pdf_pages(g['path']),'original_path':original.relative_to(root).as_posix() if original.exists() else '',
                        'source_ids':g['source_ids'],'source_urls':g['source_urls'],'published_dates':g['published_dates'],
                        'original_filenames':g['original_filenames'],'classification_basis':f.get('classification_basis',''),
                        'note':f.get('note',''),'path':target.relative_to(root).as_posix(),'storage':storage,
                        **({'also_placed':[p.relative_to(root).as_posix() for p in g['placed'][1:]]} if len(g['placed'])>1 else {})})
    records.sort(key=lambda r:(r['category'],r['label'],r['language'],r['name']))
    counts={c:sum(r['category']==c for r in records) for c in CATEGORIES}
    index={'security':config['security'],'name':config['name'],'generated':now(),'rule':RULE,'counts':counts,'files':records,
           'unclassified':unclassified,'conflicts':conflicts,'mops_codes':MOPS_CODES}
    if apply:
        lib=root/LIBRARY
        for c in CATEGORIES:
            d=lib/c; d.mkdir(parents=True,exist_ok=True); ph=d/PLACEHOLDER; held=any(d.glob('*.pdf'))
            if held and ph.exists(): ph.unlink()
            if not held and not ph.exists(): ph.write_text('目前無檔案；待補文件見研究根目錄 缺口清單.json。\n',encoding='utf-8')
        dump(lib/'PDF分類索引.json',index); write_xlsx(lib/'PDF分類索引.xlsx',index); write_html(lib/'PDF分類目錄.html',index); write_notes(lib/'整理說明.md',index,company)
    return {'status':'applied' if apply else 'dry_run','counts':counts,'actions':actions,'unclassified':unclassified,'conflicts':conflicts,
            'files':len(records),'note':'dry-run未建立任何檔案；確認後加 --apply' if not apply else '分類檔以硬連結共用內容；跨磁碟時才複製（storage=copy）'}

COLUMNS=('category','label','date_meaning','title','language','publisher','name','pages','original_filenames','source_urls','published_dates','classification_basis','note','sha256','path','storage')
def _cell(v): return '\n'.join(map(str,v)) if isinstance(v,list) else v
def write_xlsx(path,index):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb=openpyxl.Workbook(); ws=wb.active; ws.title='PDF分類索引'; ws.append(list(COLUMNS))
    for r in index['files']: ws.append([_cell(r.get(k)) for k in COLUMNS])
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='028090')
    ws.freeze_panes='A2'
    s=wb.create_sheet('統計'); s.append(['類別','份數'])
    for k,v in index['counts'].items(): s.append([k,v])
    if index['unclassified']:
        u=wb.create_sheet('未分類'); u.append(['sha256','original_filenames','reason'])
        for r in index['unclassified']: u.append([r['sha256'],_cell(r['original_filenames']),r['reason']])
    tmp=path.with_name(path.name+'.tmp'); wb.save(tmp); os.replace(tmp,path)

def write_html(path,index):
    e=html.escape
    cards=''.join(f'<span class="card">{e(CATEGORIES[k])} <b>{v}</b></span>' for k,v in index['counts'].items())
    rows=''.join(f'<tr><td>{e(CATEGORIES[r["category"]])}</td><td>{e(r["label"])}</td><td><a href="{e(urllib.parse.quote(r["path"].split("/",1)[1]))}">{e(r["name"])}</a></td>'
                 f'<td class="muted">{e(" / ".join(r["original_filenames"]))}</td><td>{e(str(r["pages"] or ""))}</td><td>{e(r["note"])}</td></tr>' for r in index['files'])
    gaps=''.join(f'<tr><td class="muted">{e(" / ".join(r["original_filenames"]))}</td><td class="gap">{e(r["reason"])}</td></tr>' for r in index['unclassified'])
    doc=('<!doctype html><html lang="zh-Hant"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
         f'<title>{e(index["name"])}PDF分類目錄</title><style>body{{font:16px/1.6 system-ui,"Microsoft JhengHei";max-width:1450px;margin:36px auto;padding:0 20px;color:#20364d;background:#f6f8fa}}'
         '.cards{display:flex;gap:10px;flex-wrap:wrap;margin:20px 0}.card{padding:8px 16px;background:#fff;border:1px solid #d4dce4;border-radius:7px}'
         'table{width:100%;border-collapse:collapse;background:#fff;font-size:14px}th,td{text-align:left;padding:10px;border-bottom:1px solid #dce3e9;vertical-align:top}'
         'th{background:#028090;color:#fff}a{color:#0059a0}.muted{color:#5d6d7a}.gap{color:#9a3412}</style>'
         f'<h1>{e(index["name"])} {e(index["security"])}｜PDF 分類目錄</h1><p>{len(index["files"])}份PDF（產生於 {e(index["generated"])}）。{e(index["rule"])}</p>'
         f'<div class="cards">{cards}</div><table><thead><tr><th>類別</th><th>年度／日期</th><th>PDF檔名</th><th>原始檔名</th><th>頁數</th><th>說明</th></tr></thead><tbody>{rows}</tbody></table>'
         +(f'<h2>未分類（需人工指定）</h2><table><thead><tr><th>原始檔名</th><th>原因</th></tr></thead><tbody>{gaps}</tbody></table>' if gaps else '')
         +'<p class="muted">非投資建議。</p></html>')
    path.write_text(doc,encoding='utf-8')

def write_notes(path,index,company):
    lines=['# PDF分類與命名','',f'檔名：{company}_{{代碼}}_會計年度或文件日期_文件類別_語言.pdf；外部研究採發布機構_日期_主題_語言.pdf。','',
           '年報依會計年度命名；股東會資料依會議日期命名，兩者不混淆。','','年度財報獨立分類，不把全年財報命名成單獨第四季財報；Q2／Q3季報標示含累計。','',
           '相同SHA256僅提供一份分類檔；同標題但不同SHA256全部加來源版短碼（SHA前8碼），不推定修訂先後。','',
           '原始文件（SHA命名）與資料庫連結保留；分類資料夾以硬連結提供可讀檔名入口，未重新壓縮或改寫PDF。','']
    lines+=[f'- {k}：{v}份' for k,v in index['counts'].items()]
    if index['unclassified']: lines+=['',f'未分類 {len(index["unclassified"])} 份，見 PDF分類索引.json 的 unclassified。']
    path.write_text('\n'.join(lines)+'\n',encoding='utf-8')

class Known:
    """Documents already held locally, in 分類PDF, or in a supplied Drive listing. Checked before any download."""
    def __init__(self,store,config,extra=None):
        self.ticker=str(config['ticker']); self.sha=set(); self.names={}; self.keys={}
        root=store.root
        for r in store.evidence():
            if r['status']!='full': continue
            self.sha.add(r['sha256']); self._name(original_name(r),'證據庫')
        for p in (root/'原始文件').glob('*'):
            if re.fullmatch(r'[0-9a-f]{64}',p.stem): self.sha.add(p.stem)
        for p in (root/LIBRARY).glob('*/*.pdf'): self._name(p.name,LIBRARY+'/'+p.parent.name)
        for r in previous_index(root):
            if r.get('sha256'): self.sha.add(r['sha256'])
            for n in [r.get('name'),r.get('filename'),r.get('original_filename')]+list(r.get('original_filenames',[]))+[original_name({'url':u}) for u in r.get('source_urls',[])]:
                self._name(n,'PDF分類索引')
        if isinstance(extra,dict): extra=extra.get('files',[])
        for x in extra or []:
            if isinstance(x,str): x={'title':x}
            if x.get('sha256'): self.sha.add(x['sha256'])
            for k in ('original_filename','filename','name','title'): self._name(x.get(k),'外部清單')
    def _name(self,name,where):
        if not name: return
        name=Path(name).name; self.names.setdefault(name.lower(),where)
        k=key_from_name(name,self.ticker)
        if k: self.keys.setdefault(k,(where,name))
    def has(self,filename=None,sha256=None):
        if sha256 and sha256 in self.sha: return 'same SHA256 already held'
        if not filename: return None
        name=Path(filename).name
        if name.lower() in self.names: return 'same filename in '+self.names[name.lower()]
        k=key_from_name(name,self.ticker)
        if k and k in self.keys:
            where,held=self.keys[k]; return f'same document {"/".join(k)} already held as {held} ({where})'
        return None

def dedupe_check(store,config,items,extra=None):
    known=Known(store,config,extra); have=[];missing=[]
    for x in items:
        if isinstance(x,str): x={'original_filename':x}
        name=x.get('original_filename') or x.get('filename') or original_name(x)
        reason=known.has(name,x.get('sha256'))
        (have if reason else missing).append({**x,'checked_name':name,**({'reason':reason} if reason else {})})
    return {'download':missing,'skip':have}
