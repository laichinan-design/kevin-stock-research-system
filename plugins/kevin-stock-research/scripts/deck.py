"""整合簡報: spec-driven PPTX with native (editable) charts, palette/font from deck_style, and structural QA.

Text uses a CJK font that ships with Windows/Office by default (deck_styles.json cjk_portable) so slides
display even where 思源黑體 is not installed; config.deck_style.cjk_font can switch back to Noto.
"""
from __future__ import annotations
import json, math, re, shutil, subprocess, tempfile
from datetime import date
from pathlib import Path
from visuals import styles

SLIDE_W, SLIDE_H = 13.333, 7.5
TYPES=('title','section','bullets','cards','kpi','image','chart','table','closing')
CHARTS=('column','bar','line','doughnut','pie')
DISCLAIMER='非投資建議'

def resolve_style(spec,config=None):
    st=styles(); ds={**((config or {}).get('deck_style') or {}),**(spec.get('style') or {})}
    pal=ds.get('palette'); font=ds.get('font')
    if pal not in st['palettes'] or font not in st['fonts']:
        raise ValueError('deck_style 未設定：先用 ppt-style-picker 選配色(01–10)與字型(A/B/C)，再 set-deck-style')
    f=st['fonts'][font]; p=st['palettes'][pal]
    return {'palette':pal,'font':font,'primary':p['primary'],'secondary':p['secondary'],'background':p['background'],
            'latin_title':f['latin_title'],'latin_body':f['latin_body'],'title_bold':f['title_bold'],
            'cjk':ds.get('cjk_font') or f['cjk_portable'],'neutral':st['neutral']}

# ---------- text fit (shared by build and check) ----------
def text_width_em(s):
    return sum(1.0 if ord(c)>0x2E7F else 0.3 if c==' ' else 0.58 for c in s)
def lines_needed(text,size_pt,width_in):
    per_line=max(1.0,width_in*72/size_pt)
    return sum(max(1,math.ceil(text_width_em(p)/per_line)) for p in (text.split('\n') or ['']))
def fits(text,size_pt,width_in,height_in,spacing=1.25):
    return lines_needed(text,size_pt,width_in-0.2)*size_pt*spacing/72<=height_in-0.1
def fit_size(text,width_in,height_in,start=20,minimum=12):
    size=start
    while size>minimum and not fits(text,size,width_in,height_in): size-=1
    return size

# ---------- builder ----------
class Deck:
    def __init__(self,spec,style):
        from pptx import Presentation
        from pptx.util import Inches
        self.spec=spec; self.st=style; self.prs=Presentation()
        self.prs.slide_width=Inches(SLIDE_W); self.prs.slide_height=Inches(SLIDE_H)
        self.blank=self.prs.slide_layouts[6]; self.warnings=[]
        self.set_theme_fonts()
    def rgb(self,hexs):
        from pptx.dml.color import RGBColor
        return RGBColor.from_string(hexs)
    def set_theme_fonts(self):
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        part=self.prs.slide_master.part.part_related_by(RT.THEME)
        xml=part.blob.decode('utf-8')
        def sub(tag,latin):
            nonlocal xml
            xml=re.sub(r'(<a:%s>\s*<a:latin typeface=")[^"]*(")'%tag,r'\g<1>%s\g<2>'%latin,xml,count=1)
            xml=re.sub(r'(<a:%s>.*?<a:ea typeface=")[^"]*(")'%tag,r'\g<1>%s\g<2>'%self.st['cjk'],xml,count=1,flags=re.S)
        sub('majorFont',self.st['latin_title']); sub('minorFont',self.st['latin_body'])
        part._blob=xml.encode('utf-8')
    def run_font(self,run,size,bold=False,color=None,title=False):
        from pptx.util import Pt
        from pptx.oxml.ns import qn
        f=run.font; f.size=Pt(size); f.bold=bold; f.name=self.st['latin_title'] if title else self.st['latin_body']
        if color: f.color.rgb=self.rgb(color)
        rpr=run._r.get_or_add_rPr()
        for tag in ('a:ea','a:cs'):
            el=rpr.find(qn(tag))
            if el is None:
                el=rpr.makeelement(qn(tag),{}); rpr.append(el)
            el.set('typeface',self.st['cjk'])
    def textbox(self,slide,x,y,w,h,paragraphs,size=18,color=None,bold=False,title=False,align=None,name=None,autofit=True,bullet=False):
        """paragraphs: list of str or (str, level)."""
        from pptx.util import Inches, Pt
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame; tf.word_wrap=True
        tf.margin_left=tf.margin_right=Inches(0.08); tf.margin_top=tf.margin_bottom=Inches(0.04); tf.vertical_anchor=MSO_ANCHOR.TOP
        if name: tb.name=name
        items=[p if isinstance(p,tuple) else (p,0) for p in paragraphs]
        joined='\n'.join(('• ' if bullet else '')+t for t,_ in items)
        if autofit:
            size=fit_size(joined,w,h,size,max(10,size-8))
            if not fits(joined,size,w,h): self.warnings.append(f'可能溢字：{joined[:24]}…')
        for i,(t,level) in enumerate(items):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.level=min(level,4)
            if align: p.alignment={'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
            if bullet: p.space_after=Pt(6)
            r=p.add_run(); r.text=(('• ' if level==0 else '– ') if bullet else '')+t
            self.run_font(r,size-2*level,bold,color or self.st['neutral']['ink'],title)
        return tb
    def rect(self,slide,x,y,w,h,fill,line=None,rounded=False):
        from pptx.util import Inches, Pt
        from pptx.enum.shapes import MSO_SHAPE
        s=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
        s.fill.solid(); s.fill.fore_color.rgb=self.rgb(fill)
        if line: s.line.color.rgb=self.rgb(line); s.line.width=Pt(1.25)
        else: s.line.fill.background()
        if rounded: s.adjustments[0]=0.08
        s.shadow.inherit=False
        return s
    def chrome(self,slide,title,source,n,total):
        N=self.st['neutral']
        self.rect(slide,0,0,SLIDE_W,0.12,self.st['primary'])
        self.rect(slide,0.5,0.42,0.09,0.62,self.st['secondary'])
        self.textbox(slide,0.72,0.3,SLIDE_W-1.3,0.85,[title],28,N['ink'],self.st['title_bold'],True,name='title')
        self.rect(slide,0.5,SLIDE_H-0.55,SLIDE_W-1.0,0.01,N['line'])
        src=source or self.spec.get('default_source')
        if not src: raise ValueError('每頁需資料來源（slide.source 或 spec.default_source）')
        self.textbox(slide,0.5,SLIDE_H-0.5,SLIDE_W-2.6,0.38,['來源：'+src],9,N['muted'],name='source',autofit=False)
        self.textbox(slide,SLIDE_W-2.0,SLIDE_H-0.5,1.5,0.38,[f'{n} / {total}'],9,N['muted'],align='right',name='page',autofit=False)
    def slide(self):
        return self.prs.slides.add_slide(self.blank)

    # slide types ---------------------------------------------------------
    def s_title(self,d,n,total):
        s=self.slide(); P=self.st['primary']; N=self.st['neutral']
        self.rect(s,0,0,SLIDE_W,SLIDE_H,P)
        self.rect(s,0.8,4.55,3.2,0.08,self.st['secondary'] if self.st['secondary']!=P else 'FFFFFF')
        self.textbox(s,0.8,1.7,SLIDE_W-1.6,1.9,[d['title']],44,'FFFFFF',True,True,name='title')
        if d.get('subtitle'): self.textbox(s,0.8,3.6,SLIDE_W-1.6,0.9,[d['subtitle']],20,'FFFFFF')
        meta=' ｜ '.join(x for x in (d.get('note'),self.spec.get('asof'),DISCLAIMER) if x)
        self.textbox(s,0.8,SLIDE_H-1.2,SLIDE_W-1.6,0.6,[meta],12,'FFFFFF',name='source',autofit=False)
    def s_section(self,d,n,total):
        s=self.slide(); self.rect(s,0,0,SLIDE_W,SLIDE_H,self.st['neutral']['pale'])
        self.rect(s,0,0,0.35,SLIDE_H,self.st['primary'])
        self.textbox(s,1.0,2.6,SLIDE_W-2,1.3,[d['title']],40,self.st['primary'],True,True,name='title')
        if d.get('subtitle'): self.textbox(s,1.0,3.9,SLIDE_W-2,0.9,[d['subtitle']],18,self.st['neutral']['muted'])
        self.textbox(s,SLIDE_W-2.0,SLIDE_H-0.5,1.5,0.38,[f'{n} / {total}'],9,self.st['neutral']['muted'],align='right',name='page',autofit=False)
    def s_bullets(self,d,n,total):
        s=self.slide(); self.chrome(s,d['title'],d.get('source'),n,total)
        items=[]
        for b in d['bullets']:
            if isinstance(b,list): items+= [(x,1) for x in b]
            else: items.append((b,0))
        w=SLIDE_W-1.2 if not d.get('takeaway') else 8.2
        self.textbox(s,0.6,1.35,w,5.4,items,22,bullet=True)
        if d.get('takeaway'): self.takeaway(s,d['takeaway'])
    def takeaway(self,s,text):
        x=9.1;self.rect(s,x,1.45,3.73,5.2,self.st['neutral']['pale'],rounded=True)
        self.textbox(s,x+0.2,1.6,3.33,0.5,['重點'],16,self.st['primary'],True,True)
        self.textbox(s,x+0.2,2.15,3.33,4.3,[text],16)
    def s_cards(self,d,n,total):
        s=self.slide(); self.chrome(s,d['title'],d.get('source'),n,total)
        cards=d['cards']; k=len(cards)
        if not 1<=k<=4: raise ValueError('cards 需 1–4 張')
        gap=0.3; w=(SLIDE_W-1.2-gap*(k-1))/k
        for i,c in enumerate(cards):
            x=0.6+i*(w+gap)
            self.rect(s,x,1.45,w,5.2,self.st['neutral']['pale'],self.st['primary'] if i==0 else self.st['neutral']['line'],True)
            self.rect(s,x,1.45,w,0.1,self.st['primary'] if i%2==0 else self.st['secondary'])
            self.textbox(s,x+0.2,1.7,w-0.4,0.9,[c['title']],20,self.st['primary'],True,True)
            self.textbox(s,x+0.2,2.6,w-0.4,3.9,c['body'] if isinstance(c['body'],list) else [c['body']],15,bullet=isinstance(c['body'],list))
    def s_kpi(self,d,n,total):
        s=self.slide(); self.chrome(s,d['title'],d.get('source'),n,total)
        stats=d['stats']; k=len(stats)
        if not 1<=k<=4: raise ValueError('kpi 需 1–4 個')
        gap=0.3; w=(SLIDE_W-1.2-gap*(k-1))/k
        for i,st in enumerate(stats):
            x=0.6+i*(w+gap); self.rect(s,x,1.6,w,2.6,'FFFFFF',self.st['neutral']['line'],True)
            self.textbox(s,x+0.15,1.8,w-0.3,1.2,[st['value']],40,self.st['primary'],True,True,align='center')
            self.textbox(s,x+0.15,3.0,w-0.3,0.5,[st['label']],16,self.st['neutral']['ink'],True,align='center')
            if st.get('note'): self.textbox(s,x+0.15,3.5,w-0.3,0.6,[st['note']],12,self.st['neutral']['muted'],align='center')
        if d.get('bullets'): self.textbox(s,0.6,4.5,SLIDE_W-1.2,2.2,d['bullets'],18,bullet=True)
    def s_image(self,d,n,total,base):
        from pptx.util import Inches
        s=self.slide(); self.chrome(s,d['title'],d.get('source'),n,total)
        p=Path(d['image']); p=p if p.is_absolute() else Path(base)/p
        if not p.exists(): raise ValueError('圖片不存在：'+str(p))
        if p.suffix.lower()=='.svg': raise ValueError('簡報請引用 PNG預覽（PowerPoint 舊版不支援 SVG）')
        from PIL import Image
        with Image.open(p) as im: iw,ih=im.size
        maxw,maxh=SLIDE_W-1.2,(4.9 if d.get('caption') else 5.3)
        sc=min(maxw/iw,maxh/ih); w,h=iw*sc,ih*sc
        s.shapes.add_picture(str(p),Inches((SLIDE_W-w)/2),Inches(1.3),Inches(w),Inches(h))
        if d.get('caption'): self.textbox(s,0.6,1.3+h+0.08,SLIDE_W-1.2,0.5,[d['caption']],13,self.st['neutral']['muted'],align='center')
    def s_chart(self,d,n,total):
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
        from pptx.util import Inches, Pt
        s=self.slide(); self.chrome(s,d['title'],d.get('source'),n,total)
        kind=d['chart']
        if kind not in CHARTS: raise ValueError('chart 須為 '+'/'.join(CHARTS))
        cats=d['categories']; cd=CategoryChartData(); cd.categories=cats
        for se in d['series']:
            if len(se['values'])!=len(cats): raise ValueError('series 長度與 categories 不符：'+se['name'])
            cd.add_series(se['name'],se['values'],number_format=d.get('number_format'))
        xl={'column':XL_CHART_TYPE.COLUMN_CLUSTERED,'bar':XL_CHART_TYPE.BAR_CLUSTERED,'line':XL_CHART_TYPE.LINE_MARKERS,
            'doughnut':XL_CHART_TYPE.DOUGHNUT,'pie':XL_CHART_TYPE.PIE}[kind]
        w=SLIDE_W-1.2 if not d.get('takeaway') else 8.2
        ch=s.shapes.add_chart(xl,Inches(0.6),Inches(1.35),Inches(w),Inches(5.4),cd).chart
        ch.font.size=Pt(12); ch.font.name=self.st['latin_body']; ch.has_title=bool(d.get('chart_title'))
        if d.get('chart_title'): ch.chart_title.text_frame.text=d['chart_title']
        if kind in ('column','bar'):
            # Bars must start at zero; a truncated axis exaggerates differences.
            ch.value_axis.minimum_scale=0; ch.value_axis.has_major_gridlines=True
        palette=[self.st['primary'],self.st['secondary'],self.st['neutral']['line'],self.st['neutral']['muted'],self.st['background']]
        if kind in ('doughnut','pie'):
            for i,pt in enumerate(ch.plots[0].series[0].points):
                pt.format.fill.solid(); pt.format.fill.fore_color.rgb=self.rgb(palette[i%len(palette)])
            ch.has_legend=True; ch.legend.position=XL_LEGEND_POSITION.RIGHT; ch.legend.include_in_layout=False
        else:
            for i,se in enumerate(ch.plots[0].series):
                col=self.rgb(palette[i%len(palette)])
                if kind=='line': se.format.line.color.rgb=col; se.format.line.width=Pt(2.5); se.smooth=False
                else: se.format.fill.solid(); se.format.fill.fore_color.rgb=col
            ch.has_legend=len(d['series'])>1
            if ch.has_legend: ch.legend.position=XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout=False
        if d.get('labels',True):
            pl=ch.plots[0]; pl.has_data_labels=True; dl=pl.data_labels; dl.font.size=Pt(11)
            if d.get('number_format'): dl.number_format=d['number_format']; dl.number_format_is_linked=False
            dl.show_value=True; dl.show_category_name=False
            if kind in ('doughnut','pie'): dl.show_percentage=False; dl.font.color.rgb=self.rgb('FFFFFF'); dl.font.bold=True
        if d.get('takeaway'): self.takeaway(s,d['takeaway'])
    def s_table(self,d,n,total):
        from pptx.util import Inches, Pt
        s=self.slide(); self.chrome(s,d['title'],d.get('source'),n,total)
        cols=d['columns']; rows=d['rows']
        if any(len(r)!=len(cols) for r in rows): raise ValueError('table 每列欄數須等於 columns')
        h=min(5.4,0.45*(len(rows)+1))
        tbl=s.shapes.add_table(len(rows)+1,len(cols),Inches(0.6),Inches(1.4),Inches(SLIDE_W-1.2),Inches(h)).table
        size=14 if len(rows)<=8 else 12 if len(rows)<=11 else 10
        for j,c in enumerate(cols):
            cell=tbl.cell(0,j); cell.fill.solid(); cell.fill.fore_color.rgb=self.rgb(self.st['primary'])
            cell.text_frame.text=''; r=cell.text_frame.paragraphs[0].add_run(); r.text=str(c); self.run_font(r,size,True,'FFFFFF')
        for i,row in enumerate(rows,1):
            for j,v in enumerate(row):
                cell=tbl.cell(i,j); cell.fill.solid(); cell.fill.fore_color.rgb=self.rgb('FFFFFF' if i%2 else self.st['neutral']['pale'])
                cell.text_frame.text=''; r=cell.text_frame.paragraphs[0].add_run(); r.text=str(v); self.run_font(r,size)
        if 1.4+h>6.8: self.warnings.append('表格列數多，可能超出頁面：'+d['title'])
    def s_closing(self,d,n,total):
        s=self.slide(); self.chrome(s,d['title'],d.get('source'),n,total)
        self.textbox(s,0.6,1.35,SLIDE_W-1.2,4.8,d.get('bullets',[]),22,bullet=True)
        self.rect(s,0.6,6.2,SLIDE_W-1.2,0.5,self.st['neutral']['pale'])
        self.textbox(s,0.75,6.24,SLIDE_W-1.5,0.42,[d.get('disclaimer') or '本簡報為研究整理，'+DISCLAIMER+'。'],12,self.st['neutral']['muted'],name='disclaimer',autofit=False)

def deck_filename(spec,config,slides):
    who=f"{config['name']}{config['ticker']}_" if config else ''
    topic=re.sub(r'[\\/:*?"<>|\s]+','',spec.get('topic') or '整合簡報')
    return f"{who}{topic}_{slides}頁_{(spec.get('asof') or date.today().isoformat()).replace('-','')}.pptx"

def build_deck(spec,out_dir,config=None,base=None):
    slides=spec.get('slides') or []
    if not slides: raise ValueError('deck spec requires slides')
    for i,d in enumerate(slides):
        if d.get('type') not in TYPES: raise ValueError(f'slide {i+1}: type 須為 '+'/'.join(TYPES))
        if not d.get('title'): raise ValueError(f'slide {i+1}: title required')
    if slides[-1]['type']!='closing': raise ValueError('最後一頁須為 closing（含結論、追蹤與非投資建議）')
    style=resolve_style(spec,config); deck=Deck(spec,style); total=len(slides); base=Path(base or out_dir)
    for n,d in enumerate(slides,1):
        fn=getattr(deck,'s_'+d['type'])
        fn(d,n,total,base) if d['type']=='image' else fn(d,n,total)
    out=Path(out_dir); out.mkdir(parents=True,exist_ok=True)
    path=out/deck_filename(spec,config,total); deck.prs.save(str(path))
    result={'file':str(path),'slides':total,'style':{k:style[k] for k in ('palette','font','latin_title','latin_body','cjk')},'build_warnings':deck.warnings}
    result['check']=check_deck(path)
    return result

def export_pdf(path,timeout=180):
    exe=shutil.which('soffice') or shutil.which('libreoffice')
    if not exe: return None,'skipped: LibreOffice not installed'
    with tempfile.TemporaryDirectory() as t:
        src=Path(t)/'deck.pptx'; shutil.copy2(path,src)
        try: r=subprocess.run([exe,'--headless','--convert-to','pdf','--outdir',t,str(src)],capture_output=True,text=True,timeout=timeout)
        except (OSError,subprocess.SubprocessError) as e: return None,'skipped: '+type(e).__name__
        pdf=Path(t)/'deck.pdf'
        if not pdf.exists(): return None,'skipped: LibreOffice conversion failed: '+(r.stderr or r.stdout).strip()[-160:]
        dest=Path(path).with_suffix('.pdf'); shutil.copy2(pdf,dest); return str(dest),'pdf_exported'

def check_deck(path):
    from pptx import Presentation
    from pptx.util import Emu
    prs=Presentation(str(path)); issues=[];warnings=[];fonts=set()
    slides=list(prs.slides)
    for i,s in enumerate(slides,1):
        texts=[sh.text_frame.text.strip() for sh in s.shapes if sh.has_text_frame]
        cells=[c.text for sh in s.shapes if sh.has_table for row in sh.table.rows for c in row.cells]
        odd=sorted(set(ch for t in texts+cells for ch in t if re.match('[\u2600-\u27BF\U0001F300-\U0001FAFF]',ch)))
        if odd: warnings.append(f'p{i}: 符號 {"".join(odd)} 在多數字型會變黑點或豆腐，改用文字')
        has_source=any(t.startswith('來源') or '資料來源' in t or t.startswith('Source') for t in texts)
        has_page=any(re.fullmatch(r'\d{1,3}(\s*/\s*\d{1,3})?',t) for t in texts)
        if i>1 and not (has_source or has_page): warnings.append(f'p{i}: 無來源／頁碼')
        for sh in s.shapes:
            if sh.left is not None and (sh.left<0 or sh.top<0 or sh.left+sh.width>prs.slide_width+Emu(10) or sh.top+sh.height>prs.slide_height+Emu(10)):
                issues.append(f'p{i}: 物件超出頁面（{sh.name}）')
            if not sh.has_text_frame or not sh.text_frame.text.strip(): continue
            sizes=[r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.name: fonts.add(r.font.name)
                    ea=r._r.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}ea')
                    if ea is not None: fonts.add(ea.get('typeface'))
            if sizes and sh.name not in ('source','page','disclaimer'):
                if not fits(sh.text_frame.text,max(sizes),Emu(sh.width).inches,Emu(sh.height).inches):
                    warnings.append(f'p{i}: 可能溢字「{sh.text_frame.text[:18]}…」')
        for sh in s.shapes:
            if getattr(sh,'has_chart',False) and sh.has_chart:
                ch=sh.chart
                try:
                    if ch.chart_type in (51,57) and ch.value_axis.minimum_scale not in (None,0): issues.append(f'p{i}: 長條圖數值軸未從 0 起算')
                except Exception: pass
    text_all=' '.join(sh.text_frame.text for s in slides[-1:] for sh in s.shapes if sh.has_text_frame)
    if DISCLAIMER not in text_all: issues.append('最後一頁缺「非投資建議」')
    return {'status':'failed' if issues else 'passed','slides':len(slides),'issues':issues,'warnings':warnings,'fonts':sorted(f for f in fonts if f)}

def apply_fonts(path,style,dest=None):
    """Re-font an existing PPTX (theme + every run's east-asian typeface) without touching layout."""
    from pptx import Presentation
    from pptx.oxml.ns import qn
    prs=Presentation(str(path)); d=Deck.__new__(Deck); d.st=style; d.prs=prs; d.set_theme_fonts(); n=0
    for s in prs.slides:
        for sh in s.shapes:
            frames=[sh.text_frame] if sh.has_text_frame else []
            if sh.has_table: frames+=[c.text_frame for row in sh.table.rows for c in row.cells]
            for tf in frames:
                for p in tf.paragraphs:
                    for r in p.runs:
                        rpr=r._r.get_or_add_rPr(); el=rpr.find(qn('a:ea'))
                        if el is None: el=rpr.makeelement(qn('a:ea'),{}); rpr.append(el)
                        el.set('typeface',style['cjk']); n+=1
    out=dest or path; prs.save(str(out)); return {'file':str(out),'runs_updated':n,'cjk':style['cjk']}
