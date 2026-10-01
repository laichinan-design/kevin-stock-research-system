#!/usr/bin/env python3
"""由 research.json 產出連動 Excel 模型（藍字為輸入、黑字為公式）。

用法：python3 build_excel.py <research.json> [輸出.xlsx]
需先跑 kevin_engine.py（DCF 敏感度表取引擎計算值）。
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BLUE = Font(name="Microsoft JhengHei", color="0000CC")
BLACK = Font(name="Microsoft JhengHei")
BOLD = Font(name="Microsoft JhengHei", bold=True)
HEAD = Font(name="Microsoft JhengHei", bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
KEY_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def header(ws, row, labels):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=lab)
        c.font, c.fill, c.border = HEAD, HEAD_FILL, BOX
        c.alignment = Alignment(horizontal="center")


def put(ws, ref, value, font=BLACK, fmt=None, fill=None):
    c = ws[ref]
    c.value, c.font, c.border = value, font, BOX
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build(r, out):
    k, m = r["kevin"], r["meta"]
    wb = Workbook()

    # ---------- 輸入 ----------
    ws = wb.active
    ws.title = "輸入"
    ws["A1"] = f"{m['name']}（{m['code']}）個股研究模型　資料日 {m['date']}"
    ws["A1"].font = Font(name="Microsoft JhengHei", bold=True, size=14)
    ws["A2"] = "藍字＝可修改輸入；黑字＝公式。金額：E 千元、I 億元、M/J NT$。"
    header(ws, 4, ["代號", "項目", "數值", "單位", "資料來源"])
    rows = [
        ("category", "類別", k["category"], "投資組合/CPO/潛力股", "", None),
        ("months", "累計月數 n", k.get("months", 4), "月", "年化因子 = 12/n", "0"),
        ("E", "1-n 月累計營收", k["E"], "千元", k.get("E_source", ""), "#,##0"),
        ("G", "毛利率", k["G"], "%", k.get("G_source", ""), "0.0%"),
        ("H", "淨利率（歸母）", k["H"], "%", k.get("H_source", ""), "0.0%"),
        ("I", "股本", k["I"], "億元", k.get("I_source", ""), "0.00"),
        ("W", "業主淨利占比", k.get("W", 1.0), "倍", k.get("W_source", "預設"), "0.000"),
        ("mult", "高倍率旗標", k.get("multiplier", 1.0), "×", "預設 1.0；1.5 為高倍率", "0.0"),
        ("M", "目前股價", k["M"], "NT$", k.get("M_source", ""), "#,##0.0"),
        ("J", "去年 EPS", k["J"], "NT$", k.get("J_source", ""), "0.00"),
        ("qEPS", "最新一季實際 EPS", k.get("latest_q_eps"), "NT$", k.get("latest_q_label", ""), "0.00"),
    ]
    ref = {}
    for i, (code, label, val, unit, src, fmt) in enumerate(rows, 5):
        put(ws, f"A{i}", code)
        put(ws, f"B{i}", label)
        put(ws, f"C{i}", val, BLUE, fmt)
        put(ws, f"D{i}", unit)
        put(ws, f"E{i}", src)
        ref[code] = f"輸入!$C${i}"
    widths(ws, [10, 22, 16, 18, 60])

    # ---------- Kevin 模型 ----------
    km = wb.create_sheet("Kevin模型")
    header(km, 1, ["代號", "項目", "公式說明", "數值"])
    a = f'IF({ref["category"]}="CPO",25,IF({ref["category"]}="潛力股",30,15))'
    c = f'IF({ref["category"]}="CPO",20,IF({ref["category"]}="潛力股",25,10))'
    g100 = "(D3*100)"
    spec = [
        ("N", "預期 EPS", "E×H×W×(12/n)×10÷(I×100000)×mult",
         f"={ref['E']}*{ref['H']}*{ref['W']}*(12/{ref['months']})*10/({ref['I']}*100000)*{ref['mult']}", "0.00"),
        ("O", "EPS 成長率", "(N−J)/J", f"=(D2-{ref['J']})/{ref['J']}", "0.0%"),
        ("P", "預期 PE_1", "a+(O×100)^(1/2)，負成長取負號", f"={a}+SIGN({g100})*ABS({g100})^(1/2)", "0.0"),
        ("Q", "預期 PE_2", "c+(O×100)^(2/3)，負成長取負號", f"={c}+SIGN({g100})*ABS({g100})^(2/3)", "0.0"),
        ("R", "股價 A", "N×P", "=D2*D4", "#,##0"),
        ("S", "股價 B", "N×Q", "=D2*D5", "#,##0"),
        ("T", "目標價", "(R+S)/2", "=(D6+D7)/2", "#,##0"),
        ("V", "漲幅空間", "(T−M)/M", f"=(D8-{ref['M']})/{ref['M']}", "+0.0%;-0.0%"),
        ("PE", "目前本益比", "M/N", f"=IF(D2>0,{ref['M']}/D2,\"N/A\")", "0.0"),
        ("chk", "一致性檢核", "最新一季 EPS 不得 > N",
         f'=IF(AND(ISNUMBER({ref["qEPS"]}),{ref["qEPS"]}>D2),"⚫ 資料矛盾：查股本/淨利率口徑","OK")', None),
    ]
    for i, (code, label, desc, f, fmt) in enumerate(spec, 2):
        put(km, f"A{i}", code)
        put(km, f"B{i}", label, BOLD if code == "T" else BLACK)
        put(km, f"C{i}", desc)
        put(km, f"D{i}", f, BOLD if code == "T" else BLACK, fmt, KEY_FILL if code in ("T", "V") else None)
    widths(km, [8, 16, 40, 16])

    # 淨利率 H 敏感度（公式連動）
    km["F1"], km["F1"].font = "H 淨利率敏感度 → 目標價 T（NT$）", BOLD
    put(km, "F2", "H \\ E 調整")
    e_adj = [-0.2, -0.1, 0, 0.1, 0.2]
    h_adj = [-0.03, -0.02, -0.01, 0, 0.01, 0.02, 0.03]
    for j, ea in enumerate(e_adj):
        put(km, f"{get_column_letter(7 + j)}2", ea, BLUE, "+0%;-0%;0%")
    for i, ha in enumerate(h_adj):
        row = 3 + i
        put(km, f"F{row}", f"={ref['H']}+({ha})", BLACK, "0.0%")
        for j in range(len(e_adj)):
            col = get_column_letter(7 + j)
            n = f"({ref['E']}*(1+{col}$2)*$F{row}*{ref['W']}*(12/{ref['months']})*10/({ref['I']}*100000)*{ref['mult']})"
            o100 = f"(({n}-{ref['J']})/{ref['J']}*100)"
            pe = f"(({a}+SIGN({o100})*ABS({o100})^(1/2))+({c}+SIGN({o100})*ABS({o100})^(2/3)))/2"
            put(km, f"{col}{row}", f"={n}*{pe}", BLACK, "#,##0",
                KEY_FILL if ha == 0 and e_adj[j] == 0 else None)
    for col in "FGHIJK":
        km.column_dimensions[col].width = 13

    # ---------- 雙法估值 ----------
    dv = wb.create_sheet("雙法估值")
    header(dv, 1, ["項目", "公式", "數值", "對現價"])
    M = ref["M"]
    dspec = [
        ("Forward PE", "M/N", f"={M}/Kevin模型!D2", "0.0", None),
        ("PEG", "Fwd PE ÷ (O×100)", "=IF(Kevin模型!D3>0,C2/(Kevin模型!D3*100),\"N/A\")", "0.00", None),
        ("PE Low", "Q×0.8", "=Kevin模型!D5*0.8", "0.0", None),
        ("PE Mid", "(P+Q)/2", "=(Kevin模型!D4+Kevin模型!D5)/2", "0.0", None),
        ("PE High", "P×1.2", "=Kevin模型!D4*1.2", "0.0", None),
        ("目標價 Low", "N×PE Low", "=Kevin模型!D2*C4", "#,##0", f"=(C7-{M})/{M}"),
        ("目標價 Mid", "N×PE Mid", "=Kevin模型!D2*C5", "#,##0", f"=(C8-{M})/{M}"),
        ("目標價 High", "N×PE High", "=Kevin模型!D2*C6", "#,##0", f"=(C9-{M})/{M}"),
        ("earnings explosion", "O>300% 或 PE>150 或 N≤0 → ⚫",
         f'=IF(OR(Kevin模型!D2<=0,Kevin模型!D3>3,{M}/Kevin模型!D2>150),"⚫ 雙法不適用","通過")', None, None),
        ("綜合燈號", "PEG<1 🟢；1–1.5 🟡；>1.5 🔴",
         '=IF(C10<>"通過","⚫",IF(NOT(ISNUMBER(C3)),"🔴",IF(C3<1,"🟢",IF(C3<=1.5,"🟡","🔴"))))', None, None),
    ]
    for i, (label, desc, f, fmt, up) in enumerate(dspec, 2):
        put(dv, f"A{i}", label)
        put(dv, f"B{i}", desc)
        put(dv, f"C{i}", f, BLACK, fmt)
        if up:
            put(dv, f"D{i}", up, BLACK, "+0.0%;-0.0%")
    widths(dv, [20, 32, 16, 12])

    # ---------- DCF ----------
    d = r.get("dcf")
    if d and d.get("fcf_base"):
        ds = wb.create_sheet("DCF")
        ds["A1"], ds["A1"].font = "DCF 兩階段 10 年（NT$ 百萬）", BOLD
        put(ds, "A2", "基期 FCF")
        put(ds, "B2", d["fcf_base"], BLUE, "#,##0")
        put(ds, "C2", d.get("fcf_base_source", ""))
        put(ds, "A3", "淨現金")
        put(ds, "B3", d.get("net_cash", 0), BLUE, "#,##0")
        put(ds, "C3", d.get("net_cash_source", ""))
        put(ds, "A4", "股數（百萬股）")
        put(ds, "B4", f"={ref['I']}*10", BLACK, "#,##0.0")
        header(ds, 6, ["情境", "機率", "1-5年成長", "6-10年成長", "WACC", "g"] +
               [f"Y{t}" for t in range(1, 11)] + ["PV(FCF)", "PV(TV)", "EV", "股權價值", "每股價值", "TV占比"])
        names = list(d["scenarios"])
        for i, name in enumerate(names, 7):
            sc = d["scenarios"][name]
            put(ds, f"A{i}", name, BOLD)
            for col, key, fmt in (("B", "prob", "0%"), ("C", "growth_1_5", "0.0%"), ("D", "growth_6_10", "0.0%"),
                                  ("E", "wacc", "0.0%"), ("F", "g", "0.0%")):
                put(ds, f"{col}{i}", sc[key], BLUE, fmt)
            for t in range(1, 11):
                col = get_column_letter(6 + t)
                prev = "$B$2" if t == 1 else f"{get_column_letter(5 + t)}{i}"
                gr = f"$C{i}" if t <= 5 else f"$D{i}"
                put(ds, f"{col}{i}", f"={prev}*(1+{gr})", BLACK, "#,##0")
            pv = "+".join(f"{get_column_letter(6 + t)}{i}/(1+$E{i})^{t}" for t in range(1, 11))
            put(ds, f"Q{i}", f"={pv}", BLACK, "#,##0")
            put(ds, f"R{i}", f"=P{i}*(1+F{i})/(E{i}-F{i})/(1+E{i})^10", BLACK, "#,##0")
            put(ds, f"S{i}", f"=Q{i}+R{i}", BLACK, "#,##0")
            put(ds, f"T{i}", f"=S{i}+$B$3", BLACK, "#,##0")
            put(ds, f"U{i}", f"=T{i}/$B$4", BOLD, "#,##0", KEY_FILL)
            put(ds, f"V{i}", f"=R{i}/S{i}", BLACK, "0%")
        last = 6 + len(names)
        put(ds, f"A{last + 1}", "機率加權每股價值", BOLD)
        put(ds, f"U{last + 1}", f"=SUMPRODUCT(B7:B{last},U7:U{last})", BOLD, "#,##0", KEY_FILL)
        put(ds, f"A{last + 2}", "對現價")
        put(ds, f"U{last + 2}", f"=(U{last + 1}-{M})/{M}", BLACK, "+0.0%;-0.0%")
        put(ds, f"A{last + 3}", "TV 檢核")
        put(ds, f"U{last + 3}", f'=IF(MAX(V7:V{last})>0.8,"⚠️ TV>80% EV","OK")')
        ds.column_dimensions["A"].width = 18
        ds.column_dimensions["C"].width = 14

        sens = (r.get("valuation") or {}).get("dcf", {}) or {}
        if sens.get("sensitivity"):
            s = sens["sensitivity"]
            top = last + 6
            ds[f"A{top}"], ds[f"A{top}"].font = "Base 情境 WACC × g 敏感度（每股 NT$，引擎計算值）", BOLD
            put(ds, f"A{top + 1}", "WACC \\ g")
            for j, g in enumerate(s["g"]):
                put(ds, f"{get_column_letter(2 + j)}{top + 1}", g, BOLD, "0.0%")
            for i, w in enumerate(s["wacc"]):
                put(ds, f"A{top + 2 + i}", w, BOLD, "0.0%")
                for j, v in enumerate(s["grid"][i]):
                    put(ds, f"{get_column_letter(2 + j)}{top + 2 + i}", v if v is not None else "g≥WACC",
                        BLACK, "#,##0", KEY_FILL if i == 2 and j == 2 else None)

    # ---------- 定性評分 ----------
    q = wb.create_sheet("定性評分")
    header(q, 1, ["類別", "項目", "分數", "說明"])
    row = 2
    comp = r.get("competition", {})
    for ms in comp.get("moat_sources", []):
        put(q, f"A{row}", "Moat (0-5)")
        put(q, f"B{row}", ms["source"])
        put(q, f"C{row}", ms["score"], BLUE)
        put(q, f"D{row}", ms.get("evidence", ""))
        row += 1
    start = row
    for name, ff in comp.get("five_forces", {}).items():
        put(q, f"A{row}", "Five Forces (1-5)")
        put(q, f"B{row}", name)
        put(q, f"C{row}", ff["score"], BLUE)
        put(q, f"D{row}", ff.get("note", ""))
        row += 1
    if row > start:
        put(q, f"A{row}", "Five Forces 平均", BOLD)
        put(q, f"C{row}", f"=AVERAGE(C{start}:C{row - 1})", BOLD, "0.0")
        row += 1
    b = r.get("bear", {})
    put(q, f"A{row}", "Bear case", BOLD)
    put(q, f"B{row}", "空方強度 (0-10)")
    put(q, f"C{row}", b.get("score"), BLUE, "0.0")
    put(q, f"D{row}", b.get("thesis", ""))
    widths(q, [18, 28, 10, 90])

    wb.save(out)
    return out


if __name__ == "__main__":
    src = Path(sys.argv[1])
    data = json.loads(src.read_text(encoding="utf-8"))
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_name(
        f"{data['meta']['code']}_{data['meta']['name']}_估值模型_{data['meta']['date'].replace('-', '')}.xlsx")
    print(build(data, out))
