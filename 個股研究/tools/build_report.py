#!/usr/bin/env python3
"""由 research.json（含 valuation 區塊）產出 Word 個股研究報告。

用法：python3 build_report.py <research.json> [輸出.docx]
章節順序固定：結論 → Kevin 模型 → 雙法估值 → 產業鏈 → 競爭 → DCF → Bear case → 訊號與失效條件 → 來源 → 免責。
"""
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

FONT = "Microsoft JhengHei"


def style_doc(doc):
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for name in ("Heading 1", "Heading 2", "Title"):
        s = doc.styles[name]
        s.font.name = FONT
        s.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        s.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)


def table(doc, head, rows):
    t = doc.add_table(rows=1, cols=len(head))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(head):
        t.rows[0].cells[i].text = str(h)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = "" if v is None else str(v)
    doc.add_paragraph()
    return t


def bullets(doc, items):
    for it in items:
        if it:
            doc.add_paragraph(str(it), style="List Bullet")


def pct(x):
    return "N/A" if x is None else f"{x:+.1%}"


def build(r, out):
    m, k, v = r["meta"], r["kevin"], r["valuation"]
    km, dv, dc = v["kevin"], v["dual"], v["dcf"]
    con = r.get("conclusion", {})
    doc = Document()
    style_doc(doc)
    doc.add_heading(f"{m['name']}（{m['code']}）個股研究報告", 0)
    doc.add_paragraph(f"資料日：{m['date']}　股價時間：{m.get('price_time', '')}　主題：{m.get('theme', '')}　市場：{m.get('market', '')}")

    doc.add_heading("一、核心結論", 1)
    p = doc.add_paragraph()
    run = p.add_run(f"Kevin 目標價 NT$ {km['T']:,.0f}（現價 NT$ {k['M']:,.1f}，空間 {pct(km['V'])}）；"
                    f"雙法 {dv['status']}" + (f"；DCF 機率加權 NT$ {dc['prob_weighted']:,.0f}（{pct(dc['upside'])}）" if dc else "")
                    + f"；空方強度 {r.get('bear', {}).get('score', 'N/A')}/10。")
    run.bold = True
    doc.add_paragraph(con.get("summary", ""))
    table(doc, ["訊號", "信心", "期間", "分數", "動作"],
          [[con.get("signal"), con.get("confidence"), con.get("horizon"), con.get("score"), con.get("action")]])
    table(doc, ["方法", "價值 (NT$)", "對現價"],
          [["Kevin 目標價 T", f"{km['T']:,.0f}", pct(km["V"])]]
          + ([[f"雙法 {s}", f"{dv['targets'][s]:,.0f}", pct(dv["upside"][s])] for s in ("Low", "Mid", "High")]
             if dv["status"] != "⚫" else [["雙法", "⚫ 不適用", ""]])
          + ([["DCF 機率加權", f"{dc['prob_weighted']:,.0f}", pct(dc["upside"])]] if dc else [])
          + ([["Bear case 下行目標", f"{r['bear']['downside_target']:,.0f}",
               pct((r['bear']['downside_target'] - k['M']) / k['M'])]] if r.get("bear", {}).get("downside_target") else []))

    doc.add_heading("二、Kevin 股價模型", 1)
    n = k.get("months", 4)
    table(doc, ["項目", "數值", "資料來源"], [
        ["類別", k["category"], "—"],
        [f"1-{n}月累計營收 (E)", f"{k['E']:,.0f} 千元", k.get("E_source", "")],
        ["毛利率 (G)", f"{k['G']:.1%}", k.get("G_source", "")],
        ["淨利率 (H)", f"{k['H']:.1%}", k.get("H_source", "")],
        ["股本 (I)", f"{k['I']:.2f} 億", k.get("I_source", "")],
        ["業主淨利占比 (W)", k.get("W", 1.0), k.get("W_source", "預設")],
        ["高倍率", f"×{k.get('multiplier', 1.0)}", "—"],
        ["目前股價 (M)", f"NT$ {k['M']:,.1f}", k.get("M_source", "")],
        ["去年 EPS (J)", f"{k['J']:.2f}", k.get("J_source", "")],
    ])
    table(doc, ["項目", "數值"], [
        ["預期 EPS (N)", f"{km['N']:.2f}"], ["EPS 成長率 (O)", f"{km['O']:.1%}"],
        ["PE_1 (P)", f"{km['P']:.1f}"], ["PE_2 (Q)", f"{km['Q']:.1f}"],
        ["股價 A (R)", f"{km['R']:,.0f}"], ["股價 B (S)", f"{km['S']:,.0f}"],
        ["目標價 (T)", f"NT$ {km['T']:,.0f}"], ["漲幅空間 (V)", pct(km["V"])],
        ["目前本益比", f"{km['current_PE']:.1f}" if km["current_PE"] else "N/A"],
    ])
    bullets(doc, km["warnings"])

    doc.add_heading("三、台股雙法估值", 1)
    if dv["status"] == "⚫":
        doc.add_paragraph(f"⚫ {dv['reason']}")
    else:
        table(doc, ["Fwd PE", "PEG", "PE Low/Mid/High", "綜合"],
              [[f"{dv['fwd_PE']:.1f}", f"{dv['PEG']:.2f}" if dv["PEG"] is not None else "N/A",
                " / ".join(f"{dv['pe_band'][s]:.1f}" for s in ("Low", "Mid", "High")), dv["status"]]])
        bullets(doc, dv["notes"])

    im = r.get("industry_map", {})
    doc.add_heading("四、產業鏈定位", 1)
    doc.add_paragraph(f"定位：{im.get('position', '')}")
    if im.get("chain"):
        table(doc, ["層級", "代表廠商", "說明"], [[c["layer"], c["players"], c["note"]] for c in im["chain"]])
    doc.add_paragraph("瓶頸（chokepoints）：")
    bullets(doc, im.get("chokepoints", []))
    doc.add_paragraph(f"價值遷移：{im.get('value_migration', '')}")
    doc.add_paragraph(f"集中度風險：{im.get('concentration_risk', '')}")
    if im.get("second_order_ideas"):
        doc.add_paragraph("二階受惠：")
        bullets(doc, im["second_order_ideas"])

    comp = r.get("competition", {})
    doc.add_heading("五、競爭格局與 Moat", 1)
    doc.add_paragraph(f"Moat 寬度：{comp.get('moat_width', '')}")
    if comp.get("moat_sources"):
        table(doc, ["Moat 來源", "分數 (0-5)", "證據"], [[s["source"], s["score"], s["evidence"]] for s in comp["moat_sources"]])
    ff = comp.get("five_forces", {})
    if ff:
        avg = sum(x["score"] for x in ff.values()) / len(ff)
        table(doc, ["Five Forces", "分數 (1-5，高=有利)", "說明"],
              [[name, x["score"], x["note"]] for name, x in ff.items()] + [["平均", f"{avg:.1f}", ""]])
    peers = [p for p in comp.get("peers", []) if p.get("name")]
    if peers:
        table(doc, ["同業", "代號", "營收成長", "毛利率", "Fwd PE", "備註"],
              [[p["name"], p["code"], p["revenue_growth"], p["gross_margin"], p["fwd_pe"], p["note"]] for p in peers])
    doc.add_paragraph(f"定價能力：{comp.get('pricing_power', '')}")

    doc.add_heading("六、DCF 交叉驗證", 1)
    if dc:
        d = r["dcf"]
        table(doc, ["情境", "機率", "1-5年成長", "6-10年成長", "WACC", "g", "每股價值", "TV 占 EV"],
              [[name, f"{s['prob']:.0%}", f"{d['scenarios'][name]['growth_1_5']:.1%}",
                f"{d['scenarios'][name]['growth_6_10']:.1%}", f"{d['scenarios'][name]['wacc']:.1%}",
                f"{d['scenarios'][name]['g']:.1%}", f"NT$ {s['per_share']:,.0f}", f"{s['tv_share']:.0%}"]
               for name, s in dc["scenarios"].items()])
        doc.add_paragraph(f"基期 FCF：NT$ {d['fcf_base']:,.0f} 百萬（{d.get('fcf_base_source', '')}）；"
                          f"淨現金 NT$ {d.get('net_cash', 0):,.0f} 百萬。{d.get('wacc_note', '')}")
        s = dc["sensitivity"]
        table(doc, ["WACC \\ g"] + [f"{g:.1%}" for g in s["g"]],
              [[f"{w:.1%}"] + [f"{x:,.0f}" if x is not None else "—" for x in s["grid"][i]] for i, w in enumerate(s["wacc"])])
        if dc["tv_warning"]:
            doc.add_paragraph("⚠️ 終值占 EV > 80%，DCF 對長期假設高度敏感，僅作交叉驗證。")
    else:
        doc.add_paragraph("未提供 FCF 資料，DCF 略過。")

    b = r.get("bear", {})
    doc.add_heading("七、Bear case（刻意單方面空方論點）", 1)
    doc.add_paragraph(f"空方強度：{b.get('score', 'N/A')} / 10")
    doc.add_paragraph(b.get("thesis", ""))
    if b.get("findings"):
        table(doc, ["面向", "發現"], [[f["dimension"], f["finding"]] for f in b["findings"]])
    if b.get("downside_target"):
        doc.add_paragraph(f"下行目標：NT$ {b['downside_target']:,.0f}（{b.get('downside_logic', '')}）")
    doc.add_paragraph("Thesis-killers（什麼會證明空方錯）：")
    bullets(doc, b.get("thesis_killers", []))

    doc.add_heading("八、催化劑與失效條件", 1)
    doc.add_paragraph("關鍵催化劑：")
    bullets(doc, con.get("key_catalysts", []))
    doc.add_paragraph("論點失效條件：")
    bullets(doc, con.get("invalidation", []))
    doc.add_paragraph("重跑時點：")
    bullets(doc, con.get("rerun_when", []))

    doc.add_heading("附錄：資料來源", 1)
    bullets(doc, r.get("sources", []))

    p = doc.add_paragraph()
    run = p.add_run("本報告為 Kevin 個人模型試算與研究筆記，非投資建議。")
    run.italic = True
    doc.save(out)
    return out


if __name__ == "__main__":
    src = Path(sys.argv[1])
    data = json.loads(src.read_text(encoding="utf-8"))
    if "valuation" not in data:
        sys.exit("請先執行 kevin_engine.py 產生 valuation 區塊")
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_name(
        f"{data['meta']['code']}_{data['meta']['name']}_個股研究報告_{data['meta']['date'].replace('-', '')}.docx")
    print(build(data, out))
